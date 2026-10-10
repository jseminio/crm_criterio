"""O recebido contra o contratado (pedido de Eduardo em 10/10/2026).

O CRM guarda o **contratado** (contratos e eventos). O **recebido** vem de uma planilha que o
Administrador importa pela tela, uma linha por recebimento: CNPJ ou grupo, competência, valor e data.
Aqui ficam a leitura da planilha, a parcela esperada do mês e a situação de cada cliente.

**Esperado do mês** = a parcela em bruto dos contratos que faturavam na competência, **sem** o 13 ÷ 12:
o caixa recebe a parcela, não o MRR. Em dezembro a parcela conta em dobro (a 13ª), `MES_DA_13A_PARCELA`
— hipótese a confirmar com o financeiro (em que mês a 13ª é cobrada).

**Situação**, sempre em palavra (PAD-002: nunca só a cor):
- **Em dia**: recebeu o esperado ou mais;
- **Parcial**: recebeu menos que o esperado;
- **Em aberto**: o mês foi importado e o cliente não aparece;
- **Sem registro**: o mês ainda não foi importado.

O CRM mostra; não cobra. A Critério entrega o relatório e o cliente cobra.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

__all__ = [
    "MES_DA_13A_PARCELA", "EM_DIA", "PARCIAL", "EM_ABERTO", "SEM_REGISTRO", "LinhaLida", "LinhaReconhecida",
    "Problema", "Leitura", "esperado_no_mes", "ler_planilha", "normalizar", "situacao",
]

MES_DA_13A_PARCELA = 12
"""Mês em que a 13ª parcela é cobrada (hipótese: dezembro). Nesse mês o esperado é a parcela × 2."""

EM_DIA, PARCIAL, EM_ABERTO, SEM_REGISTRO = "Em dia", "Parcial", "Em aberto", "Sem registro"
CENTAVOS = Decimal("0.01")
MESES = {"jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6, "jul": 7, "ago": 8, "set": 9, "out": 10,
         "nov": 11, "dez": 12}


def esperado_no_mes(parcela: Decimal, competencia: date) -> Decimal:
    """A parcela em bruto que o caixa espera na competência: em dobro no mês da 13ª."""
    return parcela * 2 if competencia.month == MES_DA_13A_PARCELA else parcela


def situacao(esperado: Decimal, recebido: Decimal, mes_importado: bool) -> str:
    if not mes_importado:
        return SEM_REGISTRO
    if recebido <= 0:
        return EM_ABERTO
    return EM_DIA if recebido >= esperado - CENTAVOS else PARCIAL


def normalizar(texto: object) -> str:
    """Minúsculas, sem acento e sem espaço repetido: para casar nome de grupo e cabeçalho."""
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t).strip().lower()


# ── Leitura da planilha ──────────────────────────────────────────────────────────────────────────────

COLUNAS = {
    "cnpj": ("cnpj",),
    "grupo": ("grupo", "cliente", "grupo economico", "razao social", "nome"),
    "competencia": ("competencia", "mes", "mes de competencia", "referencia"),
    "valor": ("valor", "valor recebido", "recebido", "valor pago"),
    "data": ("data", "data do recebimento", "data de recebimento", "recebido em", "data do pagamento"),
}


@dataclass(frozen=True)
class LinhaLida:
    """Uma linha da planilha já interpretada, antes de casar com o cliente."""

    numero: int
    cnpj: str | None
    grupo: str | None
    competencia: date
    valor: Decimal
    data: date | None


@dataclass(frozen=True)
class Problema:
    numero: int
    motivo: str
    trecho: str
    """O que a linha trazia, curto, para a pessoa achar a linha na planilha."""


@dataclass
class Leitura:
    linhas: list[LinhaLida] = field(default_factory=list)
    problemas: list[Problema] = field(default_factory=list)


@dataclass(frozen=True)
class LinhaReconhecida:
    linha: LinhaLida
    grupo_id: int
    empresa_id: int | None


def _celulas(conteudo: bytes, nome: str) -> list[list[object]]:
    if nome.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        livro = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        folha = livro.worksheets[0]
        return [list(l) for l in folha.iter_rows(values_only=True)]
    if nome.lower().endswith(".csv"):
        try:
            texto = conteudo.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = conteudo.decode("latin-1")
        dialeto = csv.Sniffer().sniff(texto[:2000], delimiters=";,\t") if texto.strip() else csv.excel
        return [list(l) for l in csv.reader(io.StringIO(texto), dialeto)]
    raise ValueError("a planilha precisa ser .xlsx ou .csv")


def _competencia(v: object) -> date | None:
    if isinstance(v, datetime):
        return date(v.year, v.month, 1)
    if isinstance(v, date):
        return date(v.year, v.month, 1)
    t = normalizar(v)
    if m := re.fullmatch(r"(\d{1,2})[/\-.](\d{4})", t):
        mes, ano = int(m[1]), int(m[2])
    elif m := re.fullmatch(r"(\d{4})[/\-.](\d{1,2})", t):
        ano, mes = int(m[1]), int(m[2])
    elif m := re.fullmatch(r"\d{1,2}[/\-.](\d{1,2})[/\-.](\d{4})", t):
        mes, ano = int(m[1]), int(m[2])
    elif m := re.fullmatch(r"([a-z]{3})[a-z]*[/\-. ]*(?:de )?(\d{2}|\d{4})", t):
        if m[1] not in MESES:
            return None
        mes, ano = MESES[m[1]], int(m[2]) + (2000 if len(m[2]) == 2 else 0)
    else:
        return None
    return date(ano, mes, 1) if 1 <= mes <= 12 and 2000 <= ano <= 2100 else None


def _valor(v: object) -> Decimal | None:
    if isinstance(v, (int, float, Decimal)) and not isinstance(v, bool):
        d = Decimal(str(v))
    else:
        t = str(v or "").replace("R$", "").replace(" ", "").strip()
        if not t:
            return None
        if "," in t:  # formato brasileiro: 1.234,56
            t = t.replace(".", "").replace(",", ".")
        try:
            d = Decimal(t)
        except InvalidOperation:
            return None
    return d.quantize(CENTAVOS, rounding=ROUND_HALF_UP)


def _data(v: object) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = str(v or "").strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(t, formato).date()
        except ValueError:
            continue
    return None


def ler_planilha(conteudo: bytes, nome: str) -> Leitura:
    """Acha o cabeçalho (na primeira linha que tiver competência e valor) e lê as linhas. Nada é
    descartado em silêncio: o que não dá para ler vira `Problema` com o número da linha."""
    leitura = Leitura()
    celulas = _celulas(conteudo, nome)
    cabecalho, inicio = None, 0
    for i, linha in enumerate(celulas[:20]):
        nomes = [normalizar(c) for c in linha]
        achadas = {chave: nomes.index(n) for chave, apelidos in COLUNAS.items() for n in nomes if n in apelidos}
        if "competencia" in achadas and "valor" in achadas and ("cnpj" in achadas or "grupo" in achadas):
            cabecalho, inicio = achadas, i + 1
            break
    if cabecalho is None:
        raise ValueError(
            "não achei o cabeçalho: a planilha precisa das colunas CNPJ (ou Grupo), Competência e Valor recebido"
        )

    def pegar(linha: list[object], chave: str) -> object:
        i = cabecalho.get(chave)
        return linha[i] if i is not None and i < len(linha) else None

    for n, linha in enumerate(celulas[inicio:], start=inicio + 1):
        if not any(str(c or "").strip() for c in linha):
            continue
        cnpj = re.sub(r"\D", "", str(pegar(linha, "cnpj") or "")) or None
        grupo = str(pegar(linha, "grupo") or "").strip() or None
        trecho = " · ".join(str(c) for c in (cnpj or grupo, pegar(linha, "competencia"), pegar(linha, "valor")) if c)[:120]
        if cnpj and len(cnpj) != 14:
            leitura.problemas.append(Problema(n, "CNPJ sem 14 dígitos", trecho))
            continue
        if not cnpj and not grupo:
            leitura.problemas.append(Problema(n, "sem CNPJ e sem grupo", trecho))
            continue
        competencia = _competencia(pegar(linha, "competencia"))
        if competencia is None:
            leitura.problemas.append(Problema(n, "competência ilegível (use mês/ano, como 10/2026)", trecho))
            continue
        valor = _valor(pegar(linha, "valor"))
        if valor is None or valor <= 0:
            leitura.problemas.append(Problema(n, "valor recebido vazio, zero ou ilegível", trecho))
            continue
        bruto_data = pegar(linha, "data")
        data_recebimento = _data(bruto_data) if bruto_data not in (None, "") else None
        if bruto_data not in (None, "") and data_recebimento is None:
            leitura.problemas.append(Problema(n, "data do recebimento ilegível (use dd/mm/aaaa)", trecho))
            continue
        leitura.linhas.append(LinhaLida(n, cnpj, grupo, competencia, valor, data_recebimento))
    return leitura
