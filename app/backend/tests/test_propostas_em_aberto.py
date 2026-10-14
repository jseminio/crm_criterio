"""Propostas em aberto · visão gerencial (amostra aprovada por Eduardo em 14/10/2026): mensalizado com a 13ª,
temperatura que soma a linha, conversão do ano e de 12 meses, enviadas no mês e On hold há mais de 60 dias."""

from __future__ import annotations

import io
from datetime import date, datetime, timezone
from decimal import Decimal as D

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session, sessionmaker

from crm.api.app import criar_app
from crm.api.propostas_em_aberto import montar
from crm.db.modelos import GrupoEconomico, Oportunidade, RegistroDeAlteracao
from crm.domain.listas import LinhaServico, MotivoRecusa, Situacao, Temperatura

HOJE = date(2026, 10, 14)


@pytest.fixture
def ids(engine):
    with Session(engine) as s:
        g = GrupoEconomico(nome="Alfa")
        s.add(g)
        s.flush()
        def op(nome, situacao, **kw):
            o = Oportunidade(grupo_id=g.id, nome=nome, situacao=situacao, **kw)
            s.add(o)
            return o
        bpo = op("BPO quente", Situacao.EM_AVALIACAO, servico="BPO Financeiro", linha_servico=LinhaServico.C1,
                 preco_mensal=D("6000"), preco_anual=D("78000"), temperatura=Temperatura.QUENTE,
                 data_colocacao=date(2026, 9, 1), data_envio_proposta=date(2026, 10, 2))
        consultoria = op("Consultoria morna", Situacao.EM_AVALIACAO, servico="Consultoria", linha_servico=LinhaServico.C2,
                         preco_anual=D("24000"), temperatura=Temperatura.MORNO, data_colocacao=date(2026, 9, 10))
        parada = op("Parada", Situacao.ON_HOLD, servico="BPO Contábil", linha_servico=LinhaServico.C1, preco_mensal=D("1200"),
                    preco_anual=D("15600"), temperatura=Temperatura.FRIO, data_colocacao=date(2026, 3, 1))
        recente = op("Pausa recente", Situacao.ON_HOLD, servico="BPO Contábil", linha_servico=LinhaServico.C1,
                     preco_mensal=D("2400"), preco_anual=D("31200"), data_colocacao=date(2026, 2, 1))
        op("Enviar", Situacao.ENVIAR_PROPOSTA, servico="BPO Contábil", linha_servico=LinhaServico.C1, preco_mensal=D("3000"),
           preco_anual=D("39000"), data_colocacao=date(2026, 10, 5))
        op("Aceita", Situacao.ACEITA, servico="BPO Contábil", linha_servico=LinhaServico.C1, preco_mensal=D("1000"),
           data_colocacao=date(2026, 5, 1), data_aceite=date(2026, 6, 1))
        op("Perdida no ano", Situacao.PERDIDA, servico="Consultoria", motivo_recusa=list(MotivoRecusa)[0],
           data_colocacao=date(2026, 4, 1))
        op("Perdida antiga", Situacao.PERDIDA, servico="Consultoria", motivo_recusa=list(MotivoRecusa)[0],
           data_colocacao=date(2025, 11, 20))
        s.flush()
        # A "Pausa recente" entrou em On hold há 20 dias, pelo histórico; a "Parada" não tem registro (vale a originação).
        s.add(RegistroDeAlteracao(quando=datetime(2026, 9, 24, 12, tzinfo=timezone.utc), usuario_email="k@c.com",
                                  acao="alterou", tabela="oportunidade", registro_id=recente.id, campo="situacao",
                                  antes="Em avaliação pela empresa", depois="On hold"))
        s.commit()
        return {"bpo": bpo.id, "consultoria": consultoria.id, "parada": parada.id, "recente": recente.id}


def test_linhas_com_mensalizado_com_13a_e_temperatura_que_soma_a_linha(engine, ids):
    with Session(engine) as s:
        q = montar(s, HOJE)
    linhas = {l.etapa: l for l in q.linhas}
    assert [l.etapa for l in q.linhas] == ["Em avaliação pela empresa", "On hold", "Enviar proposta"]
    av = linhas["Em avaliação pela empresa"]
    # BPO: 6.000 × 13 ÷ 12 = 6.500; consultoria: 24.000 ÷ 12 = 2.000
    assert (av.quantas, av.anual, av.mensalizado) == (2, D("102000.00"), D("8500.00"))
    assert (av.recorrente, av.pontual, av.quente, av.morno) == (D("6500.00"), D("2000.00"), D("6500.00"), D("2000.00"))
    for l in [*q.linhas, q.total]:
        assert l.quente + l.morno + l.frio + l.sem_classificacao == l.mensalizado
    assert q.total.quantas == 5 and q.total.mensalizado == D("8500.00") + D("1300.00") + D("2600.00") + D("3250.00")


def test_conversao_do_ano_e_de_12_meses_pela_originacao(engine, ids):
    with Session(engine) as s:
        q = montar(s, HOJE)
    assert (q.conversao_do_ano.aceitas, q.conversao_do_ano.decididas, q.conversao_do_ano.percentual) == (1, 2, D("50.0"))
    assert (q.conversao_12_meses.aceitas, q.conversao_12_meses.decididas) == (1, 3)  # inclui a perdida de nov/2025


def test_on_hold_ha_mais_de_60_dias_pelo_historico_ou_pela_originacao(engine, ids):
    with Session(engine) as s:
        q = montar(s, HOJE)
    p = {x.id: x for x in q.propostas}
    assert (p[ids["parada"]].em_on_hold_desde, p[ids["parada"]].on_hold_pela_originacao) == (date(2026, 3, 1), True)
    assert (p[ids["recente"]].em_on_hold_desde, p[ids["recente"]].dias_em_on_hold) == (date(2026, 9, 24), 20)
    assert q.on_hold_antigas == 1 and q.on_hold_antigas_mensalizado == D("1300.00")


def test_enviadas_no_mes_e_o_resumo(engine, ids):
    with Session(engine) as s:
        q = montar(s, HOJE)
    assert q.enviadas_no_mes == 1 and q.enviadas_no_mes_em_aberto == D("78000")
    assert q.resumo == [
        'Do mensalizado em aberto (R$ 15.650,00/mês), R$ 6.500,00 está quente (41,5%), 100,0% dele em '
        '"Em avaliação pela empresa".',
        "On hold: 2 proposta(s), 1 há mais de 60 dias (R$ 1.300,00/mês).",
        "Enviadas no mês: 1, com R$ 78.000,00 (valor anual) ainda em aberto.",
    ]


def test_rota_e_planilha(engine, ids):
    fabrica = sessionmaker(bind=engine, expire_on_commit=False, future=True)
    with TestClient(criar_app(fabrica)) as c:
        r = c.get("/api/inteligencia/propostas-em-aberto", params={"hoje": HOJE.isoformat()})
        assert r.status_code == 200 and D(r.json()["total"]["mensalizado"]) == D("15650.00")
        x = c.get("/api/inteligencia/propostas-em-aberto/planilha", params={"hoje": HOJE.isoformat()})
    assert x.status_code == 200 and "propostas-em-aberto-20261014.xlsx" in x.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(x.content))
    quadro = wb["Quadro"]
    assert [c.value for c in quadro[3]] == ["Etapa", "Qtd", "Valor anual", "Mensalizado", "Quente", "Morno", "Frio", "Sem class."]
    assert [c.value for c in quadro[7]][:4] == ["Total em aberto", 5, 187800.0, 15650.0]
    assert wb["Propostas"].max_row == 6  # cabeçalho + 5 em aberto
