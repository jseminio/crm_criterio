"""A Grade do funil em Excel (28/09/2026)."""

from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session, sessionmaker

from crm.api.app import criar_app
from crm.db.modelos import GrupoEconomico, Oportunidade
from crm.domain.listas import Situacao, SituacaoGrupo, Temperatura
from crm.relatorios.exportacao_do_funil import COLUNAS

TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def cliente(engine):
    fabrica = sessionmaker(bind=engine, expire_on_commit=False, future=True)
    with TestClient(criar_app(fabrica)) as c:
        yield c


def _planilha(resposta):
    return openpyxl.load_workbook(io.BytesIO(resposta.content))["Funil"]


def _linhas(ws):
    return [[c.value for c in linha] for linha in ws.iter_rows(min_row=2)]


@pytest.fixture
def duas(sessao: Session):
    grupo = GrupoEconomico(nome="Alfa", situacao=SituacaoGrupo.CLIENTE)
    sessao.add(grupo)
    sessao.flush()
    sessao.add_all([
        Oportunidade(
            grupo_id=grupo.id, nome="BPO Fiscal", situacao=Situacao.ENVIAR_PROPOSTA,
            temperatura=Temperatura.QUENTE, captador="Bruno", data_colocacao=date(2026, 9, 1),
            preco_mensal=Decimal("1500.50"), preco_anual=Decimal("18006.00"),
        ),
        Oportunidade(
            grupo_id=grupo.id, nome="DP", situacao=Situacao.ACEITA,
            captador="Ana", data_colocacao=date(2026, 8, 1),
        ),
    ])
    sessao.commit()


def test_sem_oportunidade_devolve_so_o_cabecalho(cliente):
    r = cliente.get("/api/oportunidades/exportar")
    assert r.status_code == 200
    assert r.headers["content-type"] == TIPO_XLSX
    assert f'filename="funil-{date.today():%Y-%m-%d}.xlsx"' in r.headers["content-disposition"]
    ws = _planilha(r)
    assert [c.value for c in ws[1]] == COLUNAS
    assert _linhas(ws) == []


def test_traz_as_oito_colunas_da_grade_com_numero_e_data_de_verdade(cliente, duas):
    ws = _planilha(cliente.get("/api/oportunidades/exportar"))
    linhas = _linhas(ws)
    assert linhas == [
        ["Alfa", "BPO Fiscal", Situacao.ENVIAR_PROPOSTA.value, "Quente", "Bruno",
         datetime(2026, 9, 1), 1500.5, 18006.0],
        ["Alfa", "DP", Situacao.ACEITA.value, None, "Ana", datetime(2026, 8, 1), None, None],
    ]
    assert ws["G2"].number_format == "R$ #,##0.00"
    assert ws.freeze_panes == "A2"


def test_respeita_os_filtros_da_tela(cliente, duas):
    ws = _planilha(cliente.get("/api/oportunidades/exportar", params={"situacao": Situacao.ACEITA.value}))
    assert [l[1] for l in _linhas(ws)] == ["DP"]

    ws = _planilha(cliente.get("/api/oportunidades/exportar", params={"captador": "Bruno"}))
    assert [l[1] for l in _linhas(ws)] == ["BPO Fiscal"]


def test_nao_para_nas_mil_da_grade(cliente, sessao: Session):
    grupo = GrupoEconomico(nome="Grande", situacao=SituacaoGrupo.CLIENTE)
    sessao.add(grupo)
    sessao.flush()
    sessao.add_all(
        Oportunidade(grupo_id=grupo.id, nome=f"O{i}", situacao=Situacao.ENVIAR_PROPOSTA) for i in range(1005)
    )
    sessao.commit()
    assert len(_linhas(_planilha(cliente.get("/api/oportunidades/exportar")))) == 1005


def test_exportar_nao_e_lido_como_id_de_oportunidade(cliente):
    assert cliente.get("/api/oportunidades/exportar").status_code == 200
