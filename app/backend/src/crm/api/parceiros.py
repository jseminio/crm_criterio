"""A tela Parceiros (19/10/2026, amostra aprovada por Eduardo): menu próprio, no padrão do Funil comercial — Kanban e
Grade, filtros com período, indicadores com composição e exportação para Excel. As contas ficam em
`crm.domain.parceiros`; aqui só se carrega e se grava.

Quem aprova e marca pago o fee é quem tem "Parceiros: aprovar e pagar fee" — o sócio da conta, por decisão de Eduardo.
"""

from __future__ import annotations

import io
from collections import defaultdict
from collections.abc import Callable, Iterator
from datetime import date, datetime
from decimal import Decimal
from typing import Literal

import sqlalchemy as sa
from fastapi import APIRouter, Depends, HTTPException, Query, Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from crm.api.acesso import quem_fez
from crm.db.base import agora
from crm.db.modelos import FeeDeParceiro, GrupoEconomico, Oportunidade, Parceiro, Recebimento
from crm.domain import parceiros as regras
from crm.domain.listas import (
    AcordoDoParceiro,
    EtapaDoParceiro,
    MotivoDeIndicacaoRecusada,
    NivelDeParceiro,
    SituacaoDoFee,
    TipoCanal,
    TipoDePessoa,
)

__all__ = ["roteador_de_parceiros"]

CANAIS_DE_PARCEIRO = (TipoCanal.PARCEIROS, TipoCanal.ADVOGADOS)
"""Os tipos de canal cujo texto (`canal`) a conferência liga ao cadastro de Parceiros."""


# ------------------------------------------------------------------ esquemas

class ParceiroResumo(BaseModel):
    id: int
    nome: str
    tipo_pessoa: str
    nivel: str
    acordo: str
    etapa: str
    socio_responsavel: str | None
    indicacoes: int
    abertas: int
    fechadas: int
    mrr: Decimal
    pontual: Decimal


class ParceiroEntrada(BaseModel):
    nome: str = Field(min_length=1, max_length=120)
    tipo_pessoa: TipoDePessoa = TipoDePessoa.PJ
    documento: str | None = Field(default=None, max_length=20)
    tipo: str | None = Field(default=None, max_length=60)
    nivel: NivelDeParceiro = NivelDeParceiro.INDICADOR
    etapa_manual: EtapaDoParceiro = EtapaDoParceiro.MAPEADO
    acordo: AcordoDoParceiro = AcordoDoParceiro.SEM_ACORDO
    contrato_assinado_em: date | None = None
    regra_combinada: str | None = Field(default=None, max_length=400)
    combinado_por: str | None = Field(default=None, max_length=60)
    combinado_em: date | None = None
    socio_responsavel: str | None = Field(default=None, max_length=60)
    perfil_pedido: str | None = Field(default=None, max_length=400)
    contato: str | None = Field(default=None, max_length=200)
    observacao: str | None = None


class ParceiroEdicao(BaseModel):
    nome: str | None = Field(default=None, min_length=1, max_length=120)
    tipo_pessoa: TipoDePessoa | None = None
    documento: str | None = Field(default=None, max_length=20)
    tipo: str | None = Field(default=None, max_length=60)
    nivel: NivelDeParceiro | None = None
    etapa_manual: EtapaDoParceiro | None = None
    acordo: AcordoDoParceiro | None = None
    contrato_assinado_em: date | None = None
    regra_combinada: str | None = Field(default=None, max_length=400)
    combinado_por: str | None = Field(default=None, max_length=60)
    combinado_em: date | None = None
    socio_responsavel: str | None = Field(default=None, max_length=60)
    perfil_pedido: str | None = Field(default=None, max_length=400)
    contato: str | None = Field(default=None, max_length=200)
    observacao: str | None = None


class IndicacaoDoParceiro(BaseModel):
    oportunidade_id: int
    nome: str
    grupo: str | None
    grupo_id: int
    situacao: str
    recebida_em: date | None
    validada_em: date | None
    recusada: str | None
    vinculo_ate: date | None
    mensal: Decimal | None
    anual: Decimal | None
    mrr: Decimal


class FeeResposta(BaseModel):
    id: int
    parceiro_id: int
    parceiro: str
    oportunidade_id: int
    oportunidade: str
    grupo: str | None
    regra: str
    competencia: date
    base: Decimal
    calculado: Decimal
    valor_aprovado: Decimal | None
    situacao: str
    decidido_por: str | None
    decidido_em: datetime | None
    pago_em: date | None
    motivo: str | None


class ParceiroDetalhe(BaseModel):
    id: int
    nome: str
    tipo_pessoa: str
    documento: str | None
    tipo: str | None
    nivel: str
    etapa: str
    etapa_manual: str
    acordo: str
    contrato_assinado_em: date | None
    regra_combinada: str | None
    combinado_por: str | None
    combinado_em: date | None
    socio_responsavel: str | None
    perfil_pedido: str | None
    contato: str | None
    observacao: str | None
    textos_do_canal: list[str]
    indicacoes: list[IndicacaoDoParceiro]
    fees: list[FeeResposta]


class IndicadoresDosParceiros(BaseModel):
    ativos: int
    indicacoes: int
    conversao: Decimal | None
    mrr: Decimal
    sem_contrato: int
    fees: Decimal
    fees_quantos: int
    explicacoes: dict[str, str]


class ItemDaComposicaoResposta(BaseModel):
    parceiro: str
    detalhe: str
    valor: Decimal
    referencia: str | None


class EtapaEntrada(BaseModel):
    etapa: EtapaDoParceiro


class IndicacaoEntrada(BaseModel):
    oportunidade_id: int
    recebida_em: date | None = None
    validada_em: date | None = None


class IndicacaoEdicao(BaseModel):
    recebida_em: date | None = None
    validada_em: date | None = None
    recusada: MotivoDeIndicacaoRecusada | None = None


class TextoSemParceiro(BaseModel):
    texto: str
    tipo_canal: str
    oportunidades: int
    sugestao_id: int | None


class LigarTexto(BaseModel):
    texto: str = Field(min_length=1)
    parceiro_id: int


class Aprovacao(BaseModel):
    valor: Decimal = Field(ge=0)


class Recusa(BaseModel):
    motivo: str = Field(min_length=3, max_length=300)


class Pagamento(BaseModel):
    pago_em: date


class ResultadoDoCalculo(BaseModel):
    novos: int


class OportunidadeEncontrada(BaseModel):
    id: int
    nome: str
    grupo: str | None
    situacao: str
    canal: str | None
    parceiro: str | None


# ------------------------------------------------------------------ apoio

def _hoje() -> date:
    """O dia da conta de Dormente e do vínculo. Função à parte para o teste fixar a data."""
    return date.today()


def _fee(f: FeeDeParceiro) -> FeeResposta:
    return FeeResposta(
        id=f.id, parceiro_id=f.parceiro_id, parceiro=f.parceiro.nome, oportunidade_id=f.oportunidade_id,
        oportunidade=f.oportunidade.nome, grupo=f.oportunidade.grupo.nome if f.oportunidade.grupo else None,
        regra=f.regra.value, competencia=f.competencia, base=f.base, calculado=f.calculado,
        valor_aprovado=f.valor_aprovado, situacao=f.situacao.value, decidido_por=f.decidido_por,
        decidido_em=f.decidido_em, pago_em=f.pago_em, motivo=f.motivo,
    )


def _indicacoes_por_parceiro(sessao: Session) -> dict[int, list[Oportunidade]]:
    por = defaultdict(list)
    for o in sessao.scalars(sa.select(Oportunidade).where(Oportunidade.parceiro_id.is_not(None))):
        por[o.parceiro_id].append(o)
    return por


def _resumos(sessao: Session, busca: str | None, nivel: list[NivelDeParceiro] | None,
             acordo: list[AcordoDoParceiro] | None, periodo: regras.Periodo, hoje: date) -> list[regras.ResumoDoParceiro]:
    consulta = sa.select(Parceiro).order_by(Parceiro.nome)
    if busca:
        consulta = consulta.where(Parceiro.nome.ilike(f"%{busca.strip()}%"))
    if nivel:
        consulta = consulta.where(Parceiro.nivel.in_(nivel))
    if acordo:
        consulta = consulta.where(Parceiro.acordo.in_(acordo))
    por = _indicacoes_por_parceiro(sessao)
    return [regras.resumir(p, por.get(p.id, []), periodo, hoje) for p in sessao.scalars(consulta)]


def _resposta(r: regras.ResumoDoParceiro) -> ParceiroResumo:
    return ParceiroResumo(
        id=r.id, nome=r.nome, tipo_pessoa=r.tipo_pessoa, nivel=r.nivel, acordo=r.acordo, etapa=r.etapa.value,
        socio_responsavel=r.socio_responsavel, indicacoes=r.quantas, abertas=r.abertas, fechadas=r.fechadas,
        mrr=r.mrr, pontual=r.pontual,
    )


def _fees_a_aprovar(sessao: Session) -> list[FeeDeParceiro]:
    return list(sessao.scalars(
        sa.select(FeeDeParceiro).where(FeeDeParceiro.situacao == SituacaoDoFee.A_APROVAR)
        .order_by(FeeDeParceiro.competencia, FeeDeParceiro.id)
    ))


def _documento(texto: str | None) -> str | None:
    if texto is None:
        return None
    digitos = "".join(c for c in texto if c.isdigit())
    return digitos or None


def _detalhe(sessao: Session, p: Parceiro, hoje: date) -> ParceiroDetalhe:
    indicacoes = list(sessao.scalars(
        sa.select(Oportunidade).where(Oportunidade.parceiro_id == p.id)
        .order_by(Oportunidade.indicacao_recebida_em.desc().nulls_last(), Oportunidade.data_colocacao.desc().nulls_last())
    ))
    fees = sessao.scalars(sa.select(FeeDeParceiro).where(FeeDeParceiro.parceiro_id == p.id)
                          .order_by(FeeDeParceiro.competencia.desc(), FeeDeParceiro.id.desc()))
    return ParceiroDetalhe(
        id=p.id, nome=p.nome, tipo_pessoa=p.tipo_pessoa.value, documento=p.documento, tipo=p.tipo, nivel=p.nivel.value,
        etapa=regras.etapa_do_parceiro(p.etapa_manual, p.acordo, indicacoes, hoje).value,
        etapa_manual=p.etapa_manual.value, acordo=p.acordo.value, contrato_assinado_em=p.contrato_assinado_em,
        regra_combinada=p.regra_combinada, combinado_por=p.combinado_por, combinado_em=p.combinado_em,
        socio_responsavel=p.socio_responsavel, perfil_pedido=p.perfil_pedido, contato=p.contato,
        observacao=p.observacao, textos_do_canal=list(p.textos_do_canal or []),
        indicacoes=[
            IndicacaoDoParceiro(
                oportunidade_id=o.id, nome=o.nome, grupo=o.grupo.nome if o.grupo else None, grupo_id=o.grupo_id,
                situacao=o.situacao.value, recebida_em=regras.data_da_indicacao(o), validada_em=o.indicacao_validada_em,
                recusada=o.indicacao_recusada.value if o.indicacao_recusada else None, vinculo_ate=regras.vinculo_ate(o),
                mensal=o.preco_mensal, anual=o.preco_anual, mrr=regras.mrr_da_indicacao(o),
            )
            for o in indicacoes
        ],
        fees=[_fee(f) for f in fees],
    )


def _planilha(resumos: list[regras.ResumoDoParceiro]) -> bytes:
    colunas = ["Parceiro", "PF/PJ", "Nível", "Acordo", "Etapa", "Sócio", "Indicações", "Abertas", "Fechadas",
               "MRR originado", "Pontual"]
    larguras = [30, 7, 13, 20, 13, 14, 11, 9, 9, 15, 13]
    wb = Workbook()
    ws = wb.active
    ws.title = "Parceiros"
    for col, titulo in enumerate(colunas, start=1):
        c = ws.cell(row=1, column=col, value=titulo)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="F2F2F2")
        ws.column_dimensions[c.column_letter].width = larguras[col - 1]
    for n, r in enumerate(resumos, start=2):
        for col, valor in enumerate((r.nome, r.tipo_pessoa, r.nivel, r.acordo, r.etapa.value, r.socio_responsavel,
                                     r.quantas, r.abertas, r.fechadas), start=1):
            ws.cell(row=n, column=col, value=valor)
        for col, valor in ((10, r.mrr), (11, r.pontual)):
            ws.cell(row=n, column=col, value=float(valor)).number_format = "R$ #,##0.00"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{max(len(resumos) + 1, 1)}"
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


# ------------------------------------------------------------------ rotas

def roteador_de_parceiros(obter_sessao: Callable[[], Iterator[Session]]) -> APIRouter:
    r = APIRouter(tags=["parceiros"])
    Sessao = Depends(obter_sessao, scope="function")

    def filtros(
        busca: str | None = None,
        nivel: list[NivelDeParceiro] | None = Query(default=None),
        acordo: list[AcordoDoParceiro] | None = Query(default=None),
        data_tipo: Literal["indicacao", "fechamento"] = "indicacao",
        data_de: date | None = None,
        data_ate: date | None = None,
        hoje: date | None = None,
    ) -> dict:
        return {"busca": busca, "nivel": nivel, "acordo": acordo,
                "periodo": regras.Periodo(data_de, data_ate, data_tipo), "hoje": hoje or _hoje()}

    @r.get("/api/parceiros", response_model=list[ParceiroResumo])
    def listar(f: dict = Depends(filtros), sessao: Session = Sessao) -> list[ParceiroResumo]:
        return [_resposta(x) for x in _resumos(sessao, **f)]

    @r.get("/api/parceiros/exportar")
    def exportar(f: dict = Depends(filtros), sessao: Session = Sessao) -> Response:
        return Response(
            _planilha(_resumos(sessao, **f)),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="parceiros-{date.today():%Y-%m-%d}.xlsx"'},
        )

    @r.get("/api/parceiros/indicadores", response_model=IndicadoresDosParceiros)
    def ver_indicadores(f: dict = Depends(filtros), sessao: Session = Sessao) -> IndicadoresDosParceiros:
        fees = _fees_a_aprovar(sessao)
        valores = regras.indicadores(_resumos(sessao, **f), fees)
        return IndicadoresDosParceiros(**valores, fees_quantos=len(fees), explicacoes=regras.INDICADORES)

    @r.get("/api/parceiros/indicadores/composicao", response_model=list[ItemDaComposicaoResposta])
    def ver_composicao(
        indicador: Literal["ativos", "indicacoes", "conversao", "mrr", "sem_contrato", "fees"],
        f: dict = Depends(filtros), sessao: Session = Sessao,
    ) -> list[ItemDaComposicaoResposta]:
        itens = regras.composicao(indicador, _resumos(sessao, **f), _fees_a_aprovar(sessao))
        return [ItemDaComposicaoResposta(parceiro=i.parceiro, detalhe=i.detalhe, valor=i.valor, referencia=i.referencia)
                for i in itens]

    # conferência: os textos de canal antigos, sem parceiro
    @r.get("/api/parceiros/conferencia", response_model=list[TextoSemParceiro])
    def conferencia(sessao: Session = Sessao) -> list[TextoSemParceiro]:
        linhas = sessao.execute(
            sa.select(Oportunidade.canal, Oportunidade.tipo_canal, sa.func.count())
            .where(Oportunidade.parceiro_id.is_(None), Oportunidade.tipo_canal.in_(CANAIS_DE_PARCEIRO),
                   Oportunidade.canal.is_not(None), Oportunidade.canal != "")
            .group_by(Oportunidade.canal, Oportunidade.tipo_canal)
            .order_by(sa.func.count().desc(), Oportunidade.canal)
        ).all()
        sugestoes: dict[str, int] = {}
        for p in sessao.scalars(sa.select(Parceiro)):
            for t in [p.nome, *(p.textos_do_canal or [])]:
                sugestoes.setdefault(t.strip().casefold(), p.id)
        return [TextoSemParceiro(texto=t, tipo_canal=tc.value, oportunidades=n,
                                 sugestao_id=sugestoes.get(t.strip().casefold())) for t, tc, n in linhas]

    @r.post("/api/parceiros/conferencia", response_model=ParceiroDetalhe)
    def ligar_texto(corpo: LigarTexto, sessao: Session = Sessao) -> ParceiroDetalhe:
        """Liga ao parceiro todas as oportunidades sem parceiro com aquele texto de canal. A data da indicação fica a
        da originação; a validação fica em branco, para o sócio confirmar o que já tinha vínculo."""
        p = sessao.get(Parceiro, corpo.parceiro_id)
        if p is None:
            raise HTTPException(404, "Parceiro não encontrado")
        oportunidades = list(sessao.scalars(
            sa.select(Oportunidade).where(Oportunidade.parceiro_id.is_(None), Oportunidade.canal == corpo.texto,
                                          Oportunidade.tipo_canal.in_(CANAIS_DE_PARCEIRO))
        ))
        if not oportunidades:
            raise HTTPException(404, "Nenhuma oportunidade sem parceiro com esse texto de canal")
        for o in oportunidades:
            o.parceiro_id = p.id
            o.indicacao_recebida_em = o.indicacao_recebida_em or o.data_colocacao
        if corpo.texto not in (p.textos_do_canal or []):
            p.textos_do_canal = [*(p.textos_do_canal or []), corpo.texto]
        sessao.flush()
        return _detalhe(sessao, p, _hoje())

    @r.get("/api/parceiros/oportunidades", response_model=list[OportunidadeEncontrada])
    def buscar_oportunidades(busca: str = Query(min_length=2), sessao: Session = Sessao) -> list[OportunidadeEncontrada]:
        """Para ligar uma indicação: busca pelo nome da oportunidade ou do grupo, sem exigir acesso ao Funil."""
        termo = f"%{busca.strip()}%"
        linhas = sessao.execute(
            sa.select(Oportunidade, GrupoEconomico.nome, Parceiro.nome)
            .join(GrupoEconomico, GrupoEconomico.id == Oportunidade.grupo_id)
            .outerjoin(Parceiro, Parceiro.id == Oportunidade.parceiro_id)
            .where(sa.or_(Oportunidade.nome.ilike(termo), GrupoEconomico.nome.ilike(termo)))
            .order_by(Oportunidade.data_colocacao.desc().nulls_last(), Oportunidade.id.desc())
            .limit(20)
        ).all()
        return [OportunidadeEncontrada(id=o.id, nome=o.nome, grupo=g, situacao=o.situacao.value, canal=o.canal, parceiro=pn)
                for o, g, pn in linhas]

    # fees
    @r.get("/api/parceiros/fees", response_model=list[FeeResposta])
    def listar_fees(situacao: list[SituacaoDoFee] | None = Query(default=None), sessao: Session = Sessao) -> list[FeeResposta]:
        consulta = sa.select(FeeDeParceiro).order_by(FeeDeParceiro.competencia.desc(), FeeDeParceiro.id.desc())
        if situacao:
            consulta = consulta.where(FeeDeParceiro.situacao.in_(situacao))
        return [_fee(f) for f in sessao.scalars(consulta)]

    @r.post("/api/parceiros/fees/calcular", response_model=ResultadoDoCalculo)
    def calcular(sessao: Session = Sessao) -> ResultadoDoCalculo:
        """Cria, como "A aprovar", os fees que faltam nas indicações fechadas, pelos recebidos importados do Omie."""
        novos = 0
        indicacoes = sessao.scalars(sa.select(Oportunidade).where(Oportunidade.parceiro_id.is_not(None)))
        for o in indicacoes:
            if not o.situacao.ganha:
                continue
            recebidos = sessao.execute(
                sa.select(Recebimento.competencia, sa.func.sum(Recebimento.valor))
                .where(Recebimento.grupo_id == o.grupo_id, Recebimento.ativo.is_(True))
                .group_by(Recebimento.competencia)
            ).all()
            ja = set(sessao.scalars(sa.select(FeeDeParceiro.competencia).where(FeeDeParceiro.oportunidade_id == o.id)))
            for c in regras.calcular_fees(o, [(c, Decimal(v)) for c, v in recebidos], ja):
                sessao.add(FeeDeParceiro(parceiro_id=o.parceiro_id, oportunidade_id=o.id, regra=c.regra,
                                         competencia=c.competencia, base=c.base, calculado=c.calculado))
                novos += 1
        sessao.flush()
        return ResultadoDoCalculo(novos=novos)

    def _fee_ou_404(sessao: Session, fee_id: int) -> FeeDeParceiro:
        f = sessao.get(FeeDeParceiro, fee_id)
        if f is None:
            raise HTTPException(404, "Fee não encontrado")
        return f

    @r.post("/api/parceiros/fees/{fee_id}/aprovar", response_model=FeeResposta)
    def aprovar(fee_id: int, corpo: Aprovacao, sessao: Session = Sessao) -> FeeResposta:
        f = _fee_ou_404(sessao, fee_id)
        if f.situacao is not SituacaoDoFee.A_APROVAR:
            raise HTTPException(409, f"Este fee já está {f.situacao.value.lower()}")
        f.situacao, f.valor_aprovado = SituacaoDoFee.APROVADO, corpo.valor.quantize(Decimal("0.01"))
        f.decidido_por, f.decidido_em = quem_fez(None), agora()
        sessao.flush()
        return _fee(f)

    @r.post("/api/parceiros/fees/{fee_id}/recusar", response_model=FeeResposta)
    def recusar(fee_id: int, corpo: Recusa, sessao: Session = Sessao) -> FeeResposta:
        f = _fee_ou_404(sessao, fee_id)
        if f.situacao is not SituacaoDoFee.A_APROVAR:
            raise HTTPException(409, f"Este fee já está {f.situacao.value.lower()}")
        f.situacao, f.motivo = SituacaoDoFee.RECUSADO, corpo.motivo
        f.decidido_por, f.decidido_em = quem_fez(None), agora()
        sessao.flush()
        return _fee(f)

    @r.post("/api/parceiros/fees/{fee_id}/pago", response_model=FeeResposta)
    def pagar(fee_id: int, corpo: Pagamento, sessao: Session = Sessao) -> FeeResposta:
        f = _fee_ou_404(sessao, fee_id)
        if f.situacao is not SituacaoDoFee.APROVADO:
            raise HTTPException(409, "Só fee aprovado pode ser marcado pago")
        f.situacao, f.pago_em = SituacaoDoFee.PAGO, corpo.pago_em
        sessao.flush()
        return _fee(f)

    # cadastro
    def _parceiro_ou_404(sessao: Session, parceiro_id: int) -> Parceiro:
        p = sessao.get(Parceiro, parceiro_id)
        if p is None:
            raise HTTPException(404, "Parceiro não encontrado")
        return p

    def _nome_livre(sessao: Session, nome: str, menos: int | None = None) -> None:
        existente = sessao.scalar(sa.select(Parceiro).where(sa.func.lower(Parceiro.nome) == nome.strip().lower()))
        if existente is not None and existente.id != menos:
            raise HTTPException(409, f"Já existe um parceiro chamado {existente.nome}")

    @r.post("/api/parceiros", response_model=ParceiroDetalhe, status_code=201)
    def criar(corpo: ParceiroEntrada, sessao: Session = Sessao) -> ParceiroDetalhe:
        _nome_livre(sessao, corpo.nome)
        if not corpo.etapa_manual.manual:
            raise HTTPException(422, "A etapa marcada à mão só pode ser Mapeado, Qualificado ou Estratégico")
        dados = corpo.model_dump()
        dados["nome"], dados["documento"] = corpo.nome.strip(), _documento(corpo.documento)
        p = Parceiro(**dados, textos_do_canal=[])
        sessao.add(p)
        sessao.flush()
        return _detalhe(sessao, p, _hoje())

    @r.get("/api/parceiros/{parceiro_id}", response_model=ParceiroDetalhe)
    def ver(parceiro_id: int, sessao: Session = Sessao) -> ParceiroDetalhe:
        return _detalhe(sessao, _parceiro_ou_404(sessao, parceiro_id), _hoje())

    @r.patch("/api/parceiros/{parceiro_id}", response_model=ParceiroDetalhe)
    def editar(parceiro_id: int, corpo: ParceiroEdicao, sessao: Session = Sessao) -> ParceiroDetalhe:
        p = _parceiro_ou_404(sessao, parceiro_id)
        mudancas = corpo.model_dump(exclude_unset=True)
        if "nome" in mudancas:
            _nome_livre(sessao, mudancas["nome"], p.id)
            mudancas["nome"] = mudancas["nome"].strip()
        if "documento" in mudancas:
            mudancas["documento"] = _documento(mudancas["documento"])
        if "etapa_manual" in mudancas and not mudancas["etapa_manual"].manual:
            raise HTTPException(422, "A etapa marcada à mão só pode ser Mapeado, Qualificado ou Estratégico")
        for campo, valor in mudancas.items():
            if campo in {"nome", "tipo_pessoa", "nivel", "etapa_manual", "acordo"} and valor is None:
                continue
            setattr(p, campo, valor)
        sessao.flush()
        return _detalhe(sessao, p, _hoje())

    @r.post("/api/parceiros/{parceiro_id}/etapa", response_model=ParceiroDetalhe)
    def mover(parceiro_id: int, corpo: EtapaEntrada, sessao: Session = Sessao) -> ParceiroDetalhe:
        """O arrasto do Kanban. Só as etapas manuais aceitam; as outras mudam sozinhas pelos dados."""
        p = _parceiro_ou_404(sessao, parceiro_id)
        if not corpo.etapa.manual:
            raise HTTPException(
                422, f"{corpo.etapa.value} muda sozinha pelos dados: contrato assinado, indicação validada, proposta "
                     "enviada ou seis meses sem indicação. À mão, só Mapeado, Qualificado e Estratégico.")
        p.etapa_manual = corpo.etapa
        sessao.flush()
        detalhe = _detalhe(sessao, p, _hoje())
        if detalhe.etapa != corpo.etapa.value:
            raise HTTPException(
                409, f"Ficou marcado {corpo.etapa.value}, mas os dados mantêm {p.nome} em {detalhe.etapa}.")
        return detalhe

    @r.post("/api/parceiros/{parceiro_id}/indicacoes", response_model=ParceiroDetalhe)
    def ligar_indicacao(parceiro_id: int, corpo: IndicacaoEntrada, sessao: Session = Sessao) -> ParceiroDetalhe:
        p = _parceiro_ou_404(sessao, parceiro_id)
        o = sessao.get(Oportunidade, corpo.oportunidade_id)
        if o is None:
            raise HTTPException(404, "Oportunidade não encontrada")
        if o.parceiro_id is not None and o.parceiro_id != p.id:
            outro = sessao.get(Parceiro, o.parceiro_id)
            raise HTTPException(409, f"Essa oportunidade já foi indicada por {outro.nome if outro else 'outro parceiro'}")
        o.parceiro_id = p.id
        o.indicacao_recebida_em = corpo.recebida_em or o.indicacao_recebida_em or o.data_colocacao or date.today()
        o.indicacao_validada_em = corpo.validada_em
        if o.tipo_canal is None:
            o.tipo_canal, o.canal = TipoCanal.PARCEIROS, o.canal or p.nome
        sessao.flush()
        return _detalhe(sessao, p, _hoje())

    @r.patch("/api/parceiros/indicacoes/{oportunidade_id}", response_model=ParceiroDetalhe)
    def editar_indicacao(oportunidade_id: int, corpo: IndicacaoEdicao, sessao: Session = Sessao) -> ParceiroDetalhe:
        o = sessao.get(Oportunidade, oportunidade_id)
        if o is None or o.parceiro_id is None:
            raise HTTPException(404, "Indicação não encontrada")
        for campo, valor in corpo.model_dump(exclude_unset=True).items():
            setattr(o, {"recebida_em": "indicacao_recebida_em", "validada_em": "indicacao_validada_em",
                        "recusada": "indicacao_recusada"}[campo], valor)
        sessao.flush()
        return _detalhe(sessao, sessao.get(Parceiro, o.parceiro_id), _hoje())

    @r.post("/api/parceiros/indicacoes/{oportunidade_id}/desligar", response_model=ParceiroDetalhe)
    def desligar(oportunidade_id: int, sessao: Session = Sessao) -> ParceiroDetalhe:
        """Desfaz a ligação feita por engano. Indicação com fee calculado não se desliga: o fee guarda o vínculo."""
        o = sessao.get(Oportunidade, oportunidade_id)
        if o is None or o.parceiro_id is None:
            raise HTTPException(404, "Indicação não encontrada")
        if sessao.scalar(sa.select(sa.func.count()).select_from(FeeDeParceiro)
                         .where(FeeDeParceiro.oportunidade_id == o.id)):
            raise HTTPException(409, "Esta indicação já tem fee calculado. Recuse o fee em vez de desligar.")
        p = sessao.get(Parceiro, o.parceiro_id)
        o.parceiro_id = None
        o.indicacao_validada_em = o.indicacao_recusada = None
        sessao.flush()
        return _detalhe(sessao, p, _hoje())

    return r
