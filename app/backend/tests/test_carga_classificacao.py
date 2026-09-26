"""Carga da classificação: liga a unidade ao grupo pelo CNPJ, recalcula e confere contra a planilha."""

from __future__ import annotations

from datetime import date

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from crm.carga.classificacao import Relatorio, aplicar, ler
from crm.db.modelos import ClassificacaoDoGrupo, Empresa, GrupoEconomico
from crm.domain.listas import SituacaoGrupo

CNPJ = "11222333000181"


def _planilhas(tmp_path, score="=G5*0.2+H5*0.25+L5*0.12+(6-I5)*0.12+J5*0.09+(6-K5)*0.07+M5*0.15", isc=76.0):
    # notas 5,5,2,4,2,5,5: score = 1+1.25+.6+.48+.36+.28+.75... calculado abaixo pelo próprio domínio
    from crm.domain import classificacao as regra
    n = regra.Notas(receita=5, rentabilidade=5, complexidade=2, disciplina=4, risco=2, cross_sell=5, adimplencia=5, semaforo=1, churn=1)
    pontos = float(regra.score(n))
    wb = Workbook(); ws = wb.active; ws.title = "Classificação Grupo"
    ws.append(["t"]); ws.append([]); ws.append([]); ws.append(["Grupo / Empresa", "Tipo", "Nº CNPJ", "Receita", "Margem", "Horas"])
    ws.append(["▸ Grupo Alfa", "Grupo", 1, 1000, 0.5, 10, 5, 5, 2, 4, 2, 5, 5, "01 - Controlada", pontos, "A", "A1", None, None, "Sem urgência de churn — OK"])
    ws.append(["   Alfa Ltda", "Empresa"])
    isc_ws = wb.create_sheet("ISC - Índice de Saúde")
    isc_ws.append(["Grupo Alfa", 1000, 1, "A", 100, 1, 100, 1, 100, 1, 1, 1])
    isc_ws.append(["ISC — ÍNDICE DE SAÚDE", None, None, None, isc])
    rent = Workbook(); r = rent.active; r.title = "4. Clientes"
    r.append(["Razão Social", "CNPJ", "Grupo Econômico"]); r.append(["Alfa Ltda", "11.222.333/0001-81", "Grupo Alfa"])
    a, b = tmp_path / "c.xlsx", tmp_path / "r.xlsx"; wb.save(a); rent.save(b)
    return a, b


def _grupo(sessao: Session) -> GrupoEconomico:
    g = GrupoEconomico(nome="Alfa", situacao=SituacaoGrupo.CLIENTE); sessao.add(g); sessao.flush()
    sessao.add(Empresa(grupo_id=g.id, razao_social="Alfa Ltda", cnpj=CNPJ)); sessao.flush()
    return g


def test_liga_pelo_cnpj_confere_e_grava(sessao: Session, tmp_path):
    g = _grupo(sessao)
    a, b = _planilhas(tmp_path, isc=100.0)
    rel = Relatorio(); linhas = ler(a, b, sessao, rel)
    assert rel.erros == [] and len(linhas) == 1 and linhas[0].grupo_id == g.id
    aplicar(sessao, linhas, rel, date(2026, 7, 31), "teste")
    snap = sessao.scalars(select(ClassificacaoDoGrupo)).one()
    assert snap.classe == "A" and snap.rentabilidade_da_planilha and snap.churn == 1


def test_segue_a_fusao_ate_o_grupo_que_ficou(sessao: Session, tmp_path):
    g = _grupo(sessao)
    outro = GrupoEconomico(nome="Outro", situacao=SituacaoGrupo.CLIENTE); sessao.add(outro); sessao.flush()
    g.fundido_em_id = outro.id; g.situacao = SituacaoGrupo.FUNDIDO; sessao.flush()
    a, b = _planilhas(tmp_path, isc=100.0)
    rel = Relatorio(); linhas = ler(a, b, sessao, rel)
    assert linhas[0].grupo_id == outro.id


def test_divergencia_de_isc_bloqueia(sessao: Session, tmp_path):
    _grupo(sessao)
    a, b = _planilhas(tmp_path, isc=50.0)
    rel = Relatorio(); ler(a, b, sessao, rel)
    assert not rel.pode_aplicar and any("ISC" in e for e in rel.erros)


def test_cnpj_fora_do_crm_e_erro(sessao: Session, tmp_path):
    a, b = _planilhas(tmp_path, isc=100.0)
    rel = Relatorio(); ler(a, b, sessao, rel)
    assert any("não está no CRM" in e for e in rel.erros)


def test_idempotente_por_grupo_e_referencia(sessao: Session, tmp_path):
    _grupo(sessao)
    a, b = _planilhas(tmp_path, isc=100.0)
    rel = Relatorio(); linhas = ler(a, b, sessao, rel)
    aplicar(sessao, linhas, rel, date(2026, 7, 31), "t")
    rel2 = Relatorio(); aplicar(sessao, linhas, rel2, date(2026, 7, 31), "t")
    assert rel2.criadas == 0 and rel2.ja_existiam == 1


def _planilhas_de_rentabilidade(tmp_path, nota_da_planilha=2, honorario=2500, receita=2500):
    """Médio, honorário 2.500, complexidade 4, disciplina 5, risco 4: a fórmula viva dá margem 35% (nota 2);
    com a disciplina invertida o atrito cai e a margem sobe para 48% (nota 3)."""
    from crm.domain import classificacao as regra
    n = regra.Notas(receita=3, rentabilidade=nota_da_planilha, complexidade=4, disciplina=5, risco=4, cross_sell=3, adimplencia=5, semaforo=1, churn=1)
    pontos = float(regra.score(n)); letra = regra.classe(pontos)
    wb = Workbook(); ws = wb.active; ws.title = "Classificação Grupo"
    ws.append(["t"]); ws.append([]); ws.append([]); ws.append(["Grupo / Empresa"])
    eixo = {"A": "Sem urgência de churn — OK", "B": "Sem urgência de churn — OK", "C": "Sem urgência de churn — OK"}[letra]
    ws.append(["▸ Grupo Alfa", "Grupo", 1, receita, 0.5, 16, 3, nota_da_planilha, 4, 5, 4, 3, 5, "01 - Controlada", pontos, letra, f"{letra}1", None, None, eixo])
    isc_ws = wb.create_sheet("ISC - Índice de Saúde")
    isc_ws.append(["Grupo Alfa", receita, 1, letra, 100, 1, 100, 1, 100, 1, 1, 1])
    valor = {"A": 100, "B": 60, "C": 20}[letra] * 0.33 + 100 * 0.33 + 100 * 0.34
    isc_ws.append(["ISC — ÍNDICE", None, None, None, valor])
    rent = Workbook(); r = rent.active; r.title = "4. Clientes"
    r.append(["Razão Social", "CNPJ", "Grupo Econômico", "PORTE", "Honorário (R$/mês)", "Complexidade", "Disciplina", "Risco Técnico"])
    r.append(["Alfa Ltda", "11.222.333/0001-81", "Grupo Alfa", "Médio", honorario, 4, 5, 4])
    a, b = tmp_path / "c2.xlsx", tmp_path / "r2.xlsx"; wb.save(a); rent.save(b)
    return a, b


def test_recalcula_a_rentabilidade_com_a_disciplina_invertida(sessao: Session, tmp_path):
    _grupo(sessao)
    a, b = _planilhas_de_rentabilidade(tmp_path)
    rel = Relatorio(); linhas = ler(a, b, sessao, rel, recalcular=True)
    assert rel.erros == [], rel.erros
    assert linhas[0].recalculada and linhas[0].nota_da_planilha == 2 and linhas[0].notas.rentabilidade == 3
    assert round(float(linhas[0].margem), 2) == 0.48 and linhas[0].horas == 16
    assert rel.mudancas_de_nota == 1
    aplicar(sessao, linhas, rel, date(2026, 7, 31), "t", revisao=2)
    snap = sessao.scalars(select(ClassificacaoDoGrupo)).one()
    assert snap.revisao == 2 and snap.nota_rentabilidade == 3 and snap.rentabilidade_da_planilha is False


def test_revisao_convive_com_o_snapshot_anterior(sessao: Session, tmp_path):
    _grupo(sessao)
    a, b = _planilhas_de_rentabilidade(tmp_path)
    linhas = ler(a, b, sessao, Relatorio())
    aplicar(sessao, linhas, Relatorio(), date(2026, 7, 31), "t")
    rel = Relatorio(); nova = ler(a, b, sessao, rel, recalcular=True)
    aplicar(sessao, nova, rel, date(2026, 7, 31), "t", revisao=2)
    assert rel.criadas == 1 and len(sessao.scalars(select(ClassificacaoDoGrupo)).all()) == 2


def test_honorario_divergente_de_empresa_unica_usa_o_oficial(sessao: Session, tmp_path):
    _grupo(sessao)
    a, b = _planilhas_de_rentabilidade(tmp_path, honorario=9999)
    rel = Relatorio(); linhas = ler(a, b, sessao, rel, recalcular=True)
    assert rel.erros == [] and any("usei a oficial" in x for x in rel.avisos) and linhas[0].recalculada


def test_formula_viva_que_nao_reproduz_a_nota_bloqueia(sessao: Session, tmp_path):
    _grupo(sessao)
    a, b = _planilhas_de_rentabilidade(tmp_path, nota_da_planilha=5)
    rel = Relatorio(); ler(a, b, sessao, rel, recalcular=True)
    assert any("fórmula viva" in e for e in rel.erros)
