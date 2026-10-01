"""Proposta em PowerPoint (pedido de Eduardo, 01/10/2026): preço sugerido, matrizes com marcadores,
numeração PROP CCE RJ e envio. Dados e matrizes fictícios, montados aqui com python-pptx."""

from __future__ import annotations

import io
from datetime import date, timedelta
from decimal import Decimal

import pytest
import sqlalchemy as sa
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches
from sqlalchemy.orm import sessionmaker

from crm.api.app import criar_app
from crm.db.modelos import HistoricoDePreco, Oportunidade
from crm.domain.listas import Situacao, TipoDeMatriz
from crm.domain.rentabilidade import PARAMETROS_DE_RENTABILIDADE
from crm.proposta import conta
from crm.proposta.marcadores import OBRIGATORIOS, examinar, preencher
from test_questionario import RESPOSTAS_ALFA, FonteFalsa, linha

D = Decimal
ANO = date.today().year


def matriz(*textos: str, partir: bool = False) -> bytes:
    """Um .pptx com um parágrafo por texto. `partir` quebra cada texto em dois pedaços (runs), como o
    PowerPoint faz, para provar que o marcador partido também é achado."""
    p = Presentation()
    caixa = p.slides.add_slide(p.slide_layouts[6]).shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(5)).text_frame
    for i, t in enumerate(textos):
        par = caixa.paragraphs[0] if i == 0 else caixa.add_paragraph()
        pedacos = [t[: len(t) // 2], t[len(t) // 2:]] if partir else [t]
        for pedaco in pedacos:
            par.add_run().text = pedaco
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()


def texto_do_pptx(conteudo: bytes) -> list[str]:
    p = Presentation(io.BytesIO(conteudo))
    return [par.text for s in p.slides for f in s.shapes if f.has_text_frame for par in f.text_frame.paragraphs]


MATRIZ_CONTABIL = matriz(
    "PROP CCE RJ {{numero}}", "{{cliente}}", "{{tratamento}},", "{{contextualizacao}}",
    "Contábil R$ {{valor_contabil}} mês · {{horas_contabil}} horas", "DP R$ {{valor_dp}} mês · {{horas_dp}} horas",
    "Total {{horas_total}} horas · Valor Líquido: {{valor_liquido}} · Valor Bruto: {{valor_bruto}}",
    "CNPJ: {{cnpj}} · {{regime}} · {{faturamento}}/ano · {{funcionarios}} CLT + PJs · {{sistema}}",
    partir=True,
)
MATRIZ_FINANCEIRO = matriz(
    "PROP CCE RJ {{numero}}", "{{cliente}}", "{{tratamento}},", "{{contextualizacao}}",
    "{{plano_bpo}} · {{plano_plus}} · {{plano_cfo}}",
)


class TestConta:
    def test_amostra_aprovada(self):
        # Pequeno, complexidade 2, risco 1, disciplina 3: atrito 15%, 11,5 h, custo 479,41.
        s = conta.preco_sugerido(porte="Pequeno", complexidade=2, risco=1, margem_minima=D("0.6"),
                                 margem_alvo=D("0.7"), p=PARAMETROS_DE_RENTABILIDADE)
        assert (s.atrito, s.horas, s.custo, s.custo_hora) == (D("0.15"), D("11.5"), D("479.41"), D("41.69"))
        assert (s.bruto, s.liquido) == (D("2523.19"), D("2245.64"))
        assert s.disciplina == 3 and s.complexidade_informada and s.risco_informado

    def test_nota_que_falta_entra_neutra_e_diz_isso(self):
        s = conta.preco_sugerido(porte="Micro", complexidade=None, risco=None, margem_minima=D("0.6"),
                                 margem_alvo=D("0.7"), p=PARAMETROS_DE_RENTABILIDADE)
        assert (s.complexidade, s.risco, s.atrito) == (3, 3, D("0.30"))
        assert not s.complexidade_informada and not s.risco_informado

    def test_porte_fora_da_matriz_nao_sugere(self):
        assert conta.preco_sugerido(porte="Gigante", complexidade=1, risco=1, margem_minima=D("0.6"),
                                    margem_alvo=D("0.7"), p=PARAMETROS_DE_RENTABILIDADE) is None

    def test_bruto_da_nrh_e_arredondamento_ao_multiplo_de_50(self):
        assert conta.bruto_de(D("6900"), D("0.11")) == D("7750.00")  # 7.752,81
        assert conta.bruto_de(D("2250"), D("0.11")) == D("2550.00")  # 2.528,09: mais perto de 2.550
        assert conta.arredondar_50(D("25")) == D("50") and conta.arredondar_50(D("24.99")) == D("0")
        with pytest.raises(ValueError):
            conta.bruto_de(D("100"), D("1"))

    def test_formatos(self):
        assert conta.reais(D("6900")) == "6.900,00" and conta.reais(D("1234567.8")) == "1.234.567,80"
        assert conta.reais_sem_centavos_se_inteiro(D("4500")) == "4.500"
        assert conta.reais_sem_centavos_se_inteiro(D("4500.5")) == "4.500,50"
        assert conta.faturamento_por_extenso(D("4800000")) == "R$ 4,8 milhões"
        assert conta.faturamento_por_extenso(D("1000000")) == "R$ 1 milhão"
        assert conta.faturamento_por_extenso(D("350000")) == "R$ 350 mil"
        assert conta.faturamento_por_extenso(D("2500000000")) == "R$ 2,5 bilhões"


class TestMarcadores:
    def test_examina_obrigatorios_e_desconhecidos(self):
        e = examinar(matriz("{{numero}} {{cliente}} {{valr_dp}}"), TipoDeMatriz.CONTABIL)
        assert e.desconhecidos == ["valr_dp"]
        assert "tratamento" in e.faltando and "valor_dp" in e.faltando and not e.utilizavel
        assert examinar(MATRIZ_CONTABIL, TipoDeMatriz.CONTABIL).utilizavel
        assert examinar(MATRIZ_FINANCEIRO, TipoDeMatriz.FINANCEIRO).utilizavel
        assert not examinar(MATRIZ_FINANCEIRO, TipoDeMatriz.CONTABIL).utilizavel

    def test_marcador_partido_mantem_a_formatacao_do_primeiro_pedaco(self):
        p = Presentation()
        par = p.slides.add_slide(p.slide_layouts[6]).shapes.add_textbox(0, 0, Inches(4), Inches(1)).text_frame.paragraphs[0]
        for t in ("PROP CCE RJ {{nu", "mero}}", " fim"):
            par.add_run().text = t
        par.runs[0].font.bold = True
        buf = io.BytesIO()
        p.save(buf)
        saida = Presentation(io.BytesIO(preencher(buf.getvalue(), {"numero": "154.2026"})))
        runs = saida.slides[0].shapes[0].text_frame.paragraphs[0].runs
        assert "".join(r.text for r in runs) == "PROP CCE RJ 154.2026 fim"
        assert runs[0].text == "PROP CCE RJ 154.2026" and runs[0].font.bold

    def test_varias_linhas_viram_paragrafos_e_valor_nao_e_trocado_de_novo(self):
        saida = preencher(matriz("{{contextualizacao}}", "de {{cliente}}"),
                          {"contextualizacao": "Primeira.\n\nSegunda.", "cliente": "Alfa {{numero}}", "numero": "1"})
        assert texto_do_pptx(saida) == ["Primeira.", "Segunda.", "de Alfa {{numero}}"]

    def test_marcador_sem_valor_fica_a_vista(self):
        assert texto_do_pptx(preencher(matriz("{{cliente}} {{sistema}}"), {"cliente": "Alfa"})) == ["Alfa {{sistema}}"]

    def test_obrigatorios_sao_marcadores_conhecidos(self):
        from crm.proposta.marcadores import MARCADORES
        for tipo in TipoDeMatriz:
            assert set(OBRIGATORIOS[tipo]) <= set(MARCADORES)


@pytest.fixture
def fonte():
    return FonteFalsa([linha()])


@pytest.fixture
def fabrica(engine):
    return sessionmaker(bind=engine, expire_on_commit=False, future=True)


@pytest.fixture
def cliente(fabrica, fonte):
    with TestClient(criar_app(fabrica, fonte_de_questionarios=lambda: fonte)) as c:
        yield c


def _subir(cliente, tipo="Contábil", conteudo=MATRIZ_CONTABIL, por="Karine", nome="Matriz.pptx"):
    return cliente.post(f"/api/propostas/matrizes/{tipo}", params={"nome_arquivo": nome, "enviada_por": por},
                        content=conteudo)


def _oportunidade_do_questionario(cliente) -> int:
    r = cliente.post("/api/questionarios/buscar")
    assert r.status_code == 200, r.text
    return r.json()["novos"][0]["oportunidade_id"]


def _entrada(**extra):
    base = {"matriz": "Contábil", "cliente": "Exemplo Alfa", "tratamento": "Prezada Sra. Ana Souza",
            "contextualizacao": "A Exemplo Alfa atua no comércio.", "valor_contabil": "4500", "valor_dp": "2400",
            "horas_contabil": 6, "horas_dp": 3}
    return {**base, **extra}


class TestAbaDaProposta:
    def test_rascunho_a_partir_do_questionario(self, cliente):
        _subir(cliente)
        oid = _oportunidade_do_questionario(cliente)
        r = cliente.get(f"/api/oportunidades/{oid}/proposta")
        assert r.status_code == 200, r.text
        a = r.json()
        assert a["matriz_sugerida"] == "Contábil" and a["servicos"] == ["Contábil", "Fiscal"] and not a["tem_dp"]
        s = a["sugestao"]
        # Alfa: Pequeno pela régua (não confirmado), complexidade 2 e risco 2 do questionário → atrito 20%.
        assert (s["porte"], s["porte_confirmado"], D(s["atrito"]), D(s["horas"])) == ("Pequeno", False, D("0.2"), D("12"))
        assert (D(s["bruto"]), D(s["liquido"]), s["origem_da_margem"]) == (D("2632.89"), D("2343.27"), "padrão")
        rascunho = a["rascunho"]
        assert rascunho["cliente"] == "Exemplo Alfa"
        assert rascunho["tratamento"] == "Prezado(a) Sr(a). Ana Souza"  # nunca adivinha o gênero
        assert rascunho["contextualizacao"].startswith("A Exemplo Alfa busca um novo parceiro para os serviços contábeis e fiscais.\n")
        assert D(rascunho["valor_contabil"]) == D("2350")  # sem DP no escopo: o sugerido arredondado
        assert a["perfil"]["cnpj"] == "12.345.678/0001-90" and a["perfil"]["regime"] == "Lucro Presumido"
        assert a["perfil"]["funcionarios"] == "18" and a["perfil"]["sistema"] == "N/D"
        assert a["proximo_numero"] == f"154.{ANO}" and a["revisores"] == ["Eduardo", "Karine"]
        assert a["matrizes"]["Contábil"]["utilizavel"] and a["matrizes"]["Financeiro"] is None
        assert D(a["imposto"]) == D("0.11") and a["propostas"] == []

    def test_so_financeiro_sugere_a_matriz_financeiro_com_os_planos(self, fonte, cliente):
        respostas = {**RESPOSTAS_ALFA, "servicos": ["Financeiro"], "faturamento_anual": "R$ 4.800.000,00",
                     "atividade": "comércio atacadista", "filiais_localidades": "Niterói/RJ",
                     "sistemas": {"erp": {"sistema": "Omie"}}}
        fonte.linhas = [linha(respostas=respostas)]
        oid = _oportunidade_do_questionario(cliente)
        a = cliente.get(f"/api/oportunidades/{oid}/proposta").json()
        assert a["matriz_sugerida"] == "Financeiro" and a["rascunho"]["matriz"] == "Financeiro"
        assert [D(a["rascunho"][k]) for k in ("plano_bpo", "plano_plus", "plano_cfo")] == [D(5000), D(7000), D(9000)]
        assert a["perfil"]["faturamento"] == "R$ 4,8 milhões" and a["perfil"]["sistema"] == "Omie"
        assert a["rascunho"]["contextualizacao"].startswith(
            "A Exemplo Alfa atua no segmento de comércio atacadista, com operação em Niterói/RJ, e busca um novo "
            "parceiro para os serviços financeiros.")

    def test_sem_volumetria_nao_ha_preco_sugerido(self, cliente, fabrica):
        with fabrica() as s:
            from crm.db.modelos import GrupoEconomico
            from crm.domain.listas import SituacaoGrupo
            g = GrupoEconomico(nome="Grupo Vazio", situacao=SituacaoGrupo.PROSPECT)
            s.add(g)
            s.flush()
            o = Oportunidade(grupo_id=g.id, nome="Vazio", situacao=Situacao.ENVIAR_PROPOSTA)
            s.add(o)
            s.commit()
            oid = o.id
        a = cliente.get(f"/api/oportunidades/{oid}/proposta").json()
        assert a["sugestao"] is None and "volumetria" in a["sem_sugestao"]
        assert a["rascunho"]["tratamento"] == "Prezado(a) Sr(a). [NOME]" and a["tem_dp"]

    def test_oportunidade_inexistente(self, cliente):
        assert cliente.get("/api/oportunidades/999/proposta").status_code == 404


class TestGerarEEnviar:
    def test_fluxo_completo(self, cliente, fabrica):
        _subir(cliente)
        oid = _oportunidade_do_questionario(cliente)
        r = cliente.post(f"/api/oportunidades/{oid}/proposta", json=_entrada())
        assert r.status_code == 200, r.text
        p = r.json()
        assert p["numero"] == f"154.{ANO}" and (D(p["valor_liquido"]), D(p["valor_bruto"])) == (D(6900), D(7750))
        assert p["arquivo"] == f"Exemplo Alfa_Proposta BPO Contabil_154.{ANO}.pptx" and p["enviada_em"] is None

        baixado = cliente.get(f"/api/propostas/{p['id']}/pptx")
        assert baixado.status_code == 200
        assert "filename*=UTF-8''Exemplo%20Alfa_Proposta%20BPO%20Contabil" in baixado.headers["content-disposition"]
        assert texto_do_pptx(baixado.content) == [
            f"PROP CCE RJ 154.{ANO}", "Exemplo Alfa", "Prezada Sra. Ana Souza,", "A Exemplo Alfa atua no comércio.",
            "Contábil R$ 4.500 mês · 6 horas", "DP R$ 2.400 mês · 3 horas",
            "Total 9 horas · Valor Líquido: R$ 6.900,00 · Valor Bruto: R$ 7.750,00",
            "CNPJ: 12.345.678/0001-90 · Lucro Presumido · N/D/ano · 18 CLT + PJs · N/D",
        ]

        # Gerar de novo antes de enviar: mesmo número, valores novos.
        p2 = cliente.post(f"/api/oportunidades/{oid}/proposta", json=_entrada(valor_dp=None, horas_dp=None)).json()
        assert (p2["id"], p2["numero"], D(p2["valor_liquido"])) == (p["id"], f"154.{ANO}", D(4500))
        assert "DP R$ — mês · — horas" in texto_do_pptx(cliente.get(f"/api/propostas/{p['id']}/pptx").content)
        assert cliente.get(f"/api/oportunidades/{oid}/proposta").json()["rascunho"]["valor_dp"] is None

        enviada = cliente.post(f"/api/propostas/{p['id']}/enviada", json={"por": "Karine", "em": str(date.today())})
        assert enviada.status_code == 200, enviada.text
        assert (enviada.json()["enviada_por"], enviada.json()["enviada_em"]) == ("Karine", str(date.today()))
        with fabrica() as s:
            o = s.get(Oportunidade, oid)
            assert o.situacao is Situacao.EM_AVALIACAO and o.preco_mensal == D("4500.00")
            assert {"situacao", "preco_mensal"} <= set(o.campos_do_crm)
            h = s.scalars(sa.select(HistoricoDePreco).where(HistoricoDePreco.oportunidade_id == oid)).one()
            assert h.preco_mensal_novo == D("4500.00") and h.motivo == f"Proposta 154.{ANO} enviada (valor líquido)"
        assert cliente.post(f"/api/propostas/{p['id']}/enviada", json={"por": "Karine", "em": str(date.today())}).status_code == 409

        # Depois de enviada, gerar abre número novo.
        p3 = cliente.post(f"/api/oportunidades/{oid}/proposta", json=_entrada()).json()
        assert p3["numero"] == f"155.{ANO}" and p3["id"] != p["id"]
        aba = cliente.get(f"/api/oportunidades/{oid}/proposta").json()
        assert [x["numero"] for x in aba["propostas"]] == [f"155.{ANO}", f"154.{ANO}"]
        assert aba["proximo_numero"] == f"155.{ANO}"  # a 155 ainda não foi enviada: gerar regrava ela
        assert cliente.get("/api/propostas/configuracao").json()["proximo_numero"] == 156

    def test_baixar_de_novo_usa_a_matriz_da_epoca(self, cliente):
        _subir(cliente)
        oid = _oportunidade_do_questionario(cliente)
        p = cliente.post(f"/api/oportunidades/{oid}/proposta", json=_entrada()).json()
        antes = texto_do_pptx(cliente.get(f"/api/propostas/{p['id']}/pptx").content)
        _subir(cliente, conteudo=matriz("Nova {{numero}} {{cliente}} {{tratamento}} {{contextualizacao}} {{valor_contabil}} "
                                          "{{valor_dp}} {{valor_liquido}} {{valor_bruto}} {{horas_contabil}} {{horas_dp}} {{horas_total}}"))
        assert texto_do_pptx(cliente.get(f"/api/propostas/{p['id']}/pptx").content) == antes

    def test_financeiro_nao_mexe_no_preco_da_oportunidade(self, cliente, fabrica):
        _subir(cliente, "Financeiro", MATRIZ_FINANCEIRO, por="Eduardo")
        oid = _oportunidade_do_questionario(cliente)
        r = cliente.post(f"/api/oportunidades/{oid}/proposta",
                         json=_entrada(matriz="Financeiro", plano_bpo="5000", plano_plus="7000", plano_cfo="9500.5"))
        assert r.status_code == 200, r.text
        p = r.json()
        assert p["valor_liquido"] is None and p["matriz"] == "Financeiro"
        assert "R$ 5.000 · R$ 7.000 · R$ 9.500,50" in texto_do_pptx(cliente.get(f"/api/propostas/{p['id']}/pptx").content)
        cliente.post(f"/api/propostas/{p['id']}/enviada", json={"por": "Eduardo", "em": str(date.today())})
        with fabrica() as s:
            o = s.get(Oportunidade, oid)
            assert o.situacao is Situacao.EM_AVALIACAO and o.preco_mensal is None

    @pytest.mark.parametrize("extra, erro", [
        ({"valor_contabil": None}, "honorário de Contábil"),
        ({"horas_contabil": None}, "horas de consulta por ano de Contábil"),
        ({"horas_dp": None}, "horas de consulta por ano de DP"),
        ({"matriz": "Financeiro", "plano_cfo": None}, None),
    ])
    def test_campos_que_faltam(self, cliente, extra, erro):
        _subir(cliente)
        _subir(cliente, "Financeiro", MATRIZ_FINANCEIRO)
        oid = _oportunidade_do_questionario(cliente)
        r = cliente.post(f"/api/oportunidades/{oid}/proposta", json=_entrada(**extra))
        assert r.status_code == 422
        assert erro is None or erro in r.json()["detail"]
        assert cliente.get("/api/propostas/configuracao").json()["proximo_numero"] == 154  # número não queimado

    def test_sem_matriz_ou_matriz_incompleta_nao_gera(self, cliente):
        oid = _oportunidade_do_questionario(cliente)
        r = cliente.post(f"/api/oportunidades/{oid}/proposta", json=_entrada())
        assert r.status_code == 409 and "Ainda não há matriz Contábil" in r.json()["detail"]
        m = _subir(cliente, conteudo=matriz("{{numero}} {{cliente}} {{valr_dp}}")).json()
        assert not m["utilizavel"] and m["desconhecidos"] == ["valr_dp"] and "tratamento" in m["faltando"]
        r = cliente.post(f"/api/oportunidades/{oid}/proposta", json=_entrada())
        assert r.status_code == 409 and "{{valr_dp}}" in r.json()["detail"] and "{{tratamento}}" in r.json()["detail"]

    def test_envio_valida_quem_e_quando(self, cliente):
        _subir(cliente)
        oid = _oportunidade_do_questionario(cliente)
        p = cliente.post(f"/api/oportunidades/{oid}/proposta", json=_entrada()).json()
        r = cliente.post(f"/api/propostas/{p['id']}/enviada", json={"por": "Bruno", "em": str(date.today())})
        assert r.status_code == 422 and "Eduardo, Karine" in r.json()["detail"]
        amanha = str(date.today() + timedelta(days=1))
        assert cliente.post(f"/api/propostas/{p['id']}/enviada", json={"por": "Karine", "em": amanha}).status_code == 422
        assert cliente.post("/api/propostas/999/enviada", json={"por": "Karine", "em": str(date.today())}).status_code == 404

    def test_numero_ja_usado_e_pulado(self, cliente, fabrica):
        _subir(cliente)
        oid = _oportunidade_do_questionario(cliente)
        cliente.post(f"/api/oportunidades/{oid}/proposta", json=_entrada())
        from crm.db.modelos import ConfiguracaoDeProposta
        with fabrica() as s:  # alguém volta o contador por fora da tela
            s.scalars(sa.select(ConfiguracaoDeProposta)).one().proximo_numero = 154
            s.commit()
        cliente.post(f"/api/propostas/{1}/enviada", json={"por": "Karine", "em": str(date.today())})
        assert cliente.post(f"/api/oportunidades/{oid}/proposta", json=_entrada()).json()["numero"] == f"155.{ANO}"


class TestMatrizesEConfiguracao:
    def test_subir_baixar_e_listar(self, cliente):
        m = _subir(cliente, nome="Proposta BPO Full.pptx")
        assert m.status_code == 200, m.text
        assert m.json()["encontrados"] == m.json()["obrigatorios"] == len(OBRIGATORIOS[TipoDeMatriz.CONTABIL])
        assert cliente.get(f"/api/propostas/matrizes/{m.json()['id']}/pptx").content == MATRIZ_CONTABIL
        lista = cliente.get("/api/propostas/matrizes").json()
        assert lista["em_uso"]["Contábil"]["nome_arquivo"] == "Proposta BPO Full.pptx" and lista["em_uso"]["Financeiro"] is None
        valor = next(x for x in lista["marcadores"] if x["nome"] == "valor_dp")
        assert valor["obrigatorio_em"] == ["Contábil"]
        assert next(x for x in lista["marcadores"] if x["nome"] == "numero")["obrigatorio_em"] == ["Contábil", "Financeiro"]

    @pytest.mark.parametrize("kw, status", [
        ({"conteudo": b"isto nao e pptx"}, 422), ({"nome": "matriz.pdf"}, 422), ({"por": "Bruno"}, 422),
        ({"conteudo": b""}, 400), ({"tipo": "Outra"}, 422),
    ])
    def test_subida_recusada(self, cliente, kw, status):
        assert _subir(cliente, **kw).status_code == status

    def test_configuracao(self, cliente):
        c = cliente.get("/api/propostas/configuracao").json()
        assert c["proximo_numero"] == 154 and c["revisores"] == ["Eduardo", "Karine"] and D(c["imposto"]) == D("0.11")
        assert c["ultimo_usado"] is None
        novo = {"proximo_numero": 200, "revisores": [" Eduardo ", "Karine", "Bruno"], "plano_bpo": "5500",
                "plano_plus": "7000", "plano_cfo": "9000"}
        r = cliente.put("/api/propostas/configuracao", json=novo)
        assert r.status_code == 200 and r.json()["revisores"] == ["Eduardo", "Karine", "Bruno"]
        assert cliente.put("/api/propostas/configuracao", json={**novo, "revisores": ["Ana", "ana"]}).status_code == 422

    def test_proximo_numero_nao_volta_para_um_ja_usado(self, cliente):
        _subir(cliente)
        oid = _oportunidade_do_questionario(cliente)
        cliente.post(f"/api/oportunidades/{oid}/proposta", json=_entrada())
        corpo = {"proximo_numero": 154, "revisores": ["Eduardo"], "plano_bpo": "5000", "plano_plus": "7000", "plano_cfo": "9000"}
        r = cliente.put("/api/propostas/configuracao", json=corpo)
        assert r.status_code == 422 and f"154.{ANO} já foi usado" in r.json()["detail"]
        assert cliente.get("/api/propostas/configuracao").json()["ultimo_usado"] == f"154.{ANO}"
