"""Os marcadores das matrizes de proposta e o preenchimento do PowerPoint (01/10/2026).

A matriz é um .pptx comum, editado no PowerPoint, com marcadores como `{{cliente}}` no lugar do
texto que muda a cada cliente. O CRM troca cada marcador pelo valor e mantém a formatação do
trecho onde ele estava. Trocar a matriz é subir outro .pptx com os mesmos marcadores: o CRM confere
antes de aceitar (`examinar`), e matriz com marcador obrigatório faltando ou desconhecido não é usada.

O PowerPoint costuma partir o texto de um parágrafo em vários pedaços (runs) — por corretor, por
formatação —, e um marcador pode cair partido em dois. Por isso a busca é no texto inteiro do
parágrafo, e a troca junta os pedaços que o marcador ocupava.
"""

from __future__ import annotations

import copy
import io
import re
from dataclasses import dataclass
from typing import Iterator

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.text.text import _Paragraph  # interno, mas estável desde a 0.6: é o que `text_frame.paragraphs` devolve

from crm.domain.listas import TipoDeMatriz

__all__ = ["MARCADORES", "OBRIGATORIOS", "Exame", "examinar", "preencher", "MatrizIlegivel"]

PADRAO = re.compile(r"\{\{\s*([A-Za-z_]+)\s*\}\}")

#: Todos os marcadores que o CRM sabe preencher, com o que entra no lugar.
MARCADORES: dict[str, str] = {
    "numero": 'número da proposta com o ano: "154.2026" (a matriz traz "PROP CCE RJ {{numero}}")',
    "cliente": "nome do cliente na proposta",
    "tratamento": 'abertura da carta: "Prezado(a) Sr(a). Ana Souza"',
    "contextualizacao": "texto da contextualização (um parágrafo por linha, se o marcador estiver sozinho no parágrafo)",
    "valor_contabil": 'honorário mensal líquido de Contábil/Fiscal, sem "R$": "4.500"',
    "valor_dp": 'honorário mensal líquido de Departamento Pessoal, sem "R$": "2.400"',
    "horas_contabil": "horas de consulta e reuniões por ano, Contábil/Fiscal",
    "horas_dp": "horas de consulta e reuniões por ano, DP",
    "horas_total": "soma das horas de consulta por ano",
    "valor_liquido": 'total mensal líquido, com "R$": "R$ 6.900,00"',
    "valor_bruto": 'total mensal bruto com a alíquota estimada, com "R$": "R$ 7.750,00"',
    "plano_bpo": 'preço do BPO Financeiro, com "R$": "R$ 5.000"',
    "plano_plus": 'preço do BPO Financeiro PLUS, com "R$": "R$ 7.000"',
    "plano_cfo": 'preço do CFO as a Service, com "R$": "R$ 9.000"',
    "cnpj": "CNPJ principal, formatado",
    "regime": "regime tributário",
    "faturamento": 'faturamento anual informado: "R$ 4,8 milhões"',
    "funcionarios": "empregados CLT",
    "movimentacao": "lançamentos contábeis por mês",
    "volume_documentos": "documentos fiscais por mês (notas emitidas + recebidas)",
    "instituicoes": "contas bancárias",
    "meios_de_pagamento": "meios de recebimento",
    "sistema": "sistema / ERP em uso",
    "segmento": "atividade ou segmento",
    "empresas": 'empresas no escopo: "3 empresas"',
    "localidade": "cidades, UFs ou filiais",
}

_COMUNS = ("numero", "cliente", "tratamento", "contextualizacao")
#: Sem estes a proposta sai errada (com valor ou nome de outro cliente): a matriz não é usada.
OBRIGATORIOS: dict[TipoDeMatriz, tuple[str, ...]] = {
    TipoDeMatriz.CONTABIL: _COMUNS + (
        "valor_contabil", "valor_dp", "valor_liquido", "valor_bruto", "horas_contabil", "horas_dp", "horas_total",
    ),
    TipoDeMatriz.FINANCEIRO: _COMUNS + ("plano_bpo", "plano_plus", "plano_cfo"),
}


class MatrizIlegivel(ValueError):
    """O arquivo não abre como PowerPoint (.pptx)."""


@dataclass(frozen=True)
class Exame:
    encontrados: list[str]
    faltando: list[str]
    desconhecidos: list[str]

    @property
    def utilizavel(self) -> bool:
        return not self.faltando and not self.desconhecidos


def _abrir(conteudo: bytes):
    try:
        return Presentation(io.BytesIO(conteudo))
    except Exception as falha:  # python-pptx levanta tipos variados para arquivo que não é pptx
        raise MatrizIlegivel("o arquivo não abre como PowerPoint (.pptx)") from falha


def _formas(formas) -> Iterator:
    for forma in formas:
        if forma.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _formas(forma.shapes)
        else:
            yield forma


def _paragrafos(apresentacao) -> Iterator:
    for slide in apresentacao.slides:
        for forma in _formas(slide.shapes):
            if getattr(forma, "has_text_frame", False) and forma.has_text_frame:
                yield from forma.text_frame.paragraphs
            if getattr(forma, "has_table", False) and forma.has_table:
                for linha in forma.table.rows:
                    for celula in linha.cells:
                        yield from celula.text_frame.paragraphs


def _texto(paragrafo) -> str:
    return "".join(r.text for r in paragrafo.runs)


def examinar(conteudo: bytes, tipo: TipoDeMatriz) -> Exame:
    nomes: list[str] = []
    for p in _paragrafos(_abrir(conteudo)):
        for m in PADRAO.finditer(_texto(p)):
            if m.group(1) not in nomes:
                nomes.append(m.group(1))
    return Exame(
        encontrados=nomes,
        faltando=[n for n in OBRIGATORIOS[tipo] if n not in nomes],
        desconhecidos=[n for n in nomes if n not in MARCADORES],
    )


def _trocar_no_paragrafo(paragrafo, valores: dict[str, str]) -> None:
    runs = paragrafo.runs
    desde = 0  # depois do último valor posto: um valor com "{{...}}" dentro não é trocado de novo
    while True:
        textos = [r.text for r in runs]
        inteiro = "".join(textos)
        m = next((m for m in PADRAO.finditer(inteiro, desde) if m.group(1) in valores), None)
        if m is None:
            return
        inicio, fim = m.span()
        valor = valores[m.group(1)].replace("\n", " ")
        desde = inicio + len(valor)
        pos = 0
        primeiro = ultimo = None
        for i, t in enumerate(textos):
            if primeiro is None and inicio < pos + len(t):
                primeiro, ini_primeiro = i, pos
            if fim <= pos + len(t):
                ultimo, ini_ultimo = i, pos
                break
            pos += len(t)
        resto = textos[ultimo][fim - ini_ultimo:]
        runs[primeiro].text = textos[primeiro][: inicio - ini_primeiro] + valor + resto
        for i in range(primeiro + 1, ultimo + 1):
            runs[i].text = ""


def _paragrafo_em_linhas(paragrafo, nome: str, valor: str) -> bool:
    """Marcador sozinho no parágrafo e valor com várias linhas: um parágrafo por linha, com a mesma
    formatação. Devolve se tratou."""
    linhas = [l.strip() for l in valor.split("\n") if l.strip()]
    if len(linhas) < 2 or not re.fullmatch(r"\s*\{\{\s*" + nome + r"\s*\}\}\s*", _texto(paragrafo)):
        return False
    elementos = [paragrafo._p]
    for _ in linhas[1:]:
        novo = copy.deepcopy(paragrafo._p)
        elementos[-1].addnext(novo)
        elementos.append(novo)
    for elemento, linha in zip(elementos, linhas):
        _trocar_no_paragrafo(_Paragraph(elemento, paragrafo._parent), {nome: linha})
    return True


def preencher(conteudo: bytes, valores: dict[str, str]) -> bytes:
    """Troca cada marcador conhecido pelo valor. Marcador sem valor fica como está, à vista, para a
    revisão pegar — nunca some calado."""
    apresentacao = _abrir(conteudo)
    for p in list(_paragrafos(apresentacao)):
        texto = _texto(p)
        nomes = {m.group(1) for m in PADRAO.finditer(texto)}
        multilinha = [n for n in nomes if n in valores and "\n" in valores[n].strip()]
        if len(nomes) == 1 and multilinha and _paragrafo_em_linhas(p, multilinha[0], valores[multilinha[0]]):
            continue
        _trocar_no_paragrafo(p, valores)
    saida = io.BytesIO()
    apresentacao.save(saida)
    return saida.getvalue()
