"""Propostas em aberto · visão gerencial (amostra aprovada por Eduardo em 14/10/2026), para a war room com o Bruno.

- **Linhas:** as etapas em aberto (Em avaliação pela empresa, On hold, Enviar proposta), com quantidade, valor anual,
  mensalizado e o mensalizado aberto por temperatura (as colunas somam a linha).
- **Mensalizado com a 13ª** (decisão do mesmo dia): serviço recorrente (C1) pela parcela × 13 ÷ 12, como todo MRR do
  CRM; consultoria e legalização (C2) pelo valor único ÷ 12.
- **Conversão** do ano e dos últimos 12 meses pela regra do Funil (aceitas ÷ decididas), pela data de originação.
- **Enviadas no mês:** propostas com a data de envio no mês corrente, e quanto delas segue em aberto (valor anual).
- **On hold há mais de 60 dias** (troca aprovada da linha "sem data de colocação"): a entrada em On hold é a última
  mudança de situação para On hold no histórico de alterações; sem registro (carga, ou antes do histórico), vale a
  data de originação.
- **O que o quadro mostra:** frases montadas com os números do próprio quadro, sem interpretação.
"""

from __future__ import annotations

import io
from collections.abc import Callable, Iterator
from datetime import date, timedelta
from decimal import Decimal

import sqlalchemy as sa
from fastapi import APIRouter, Depends, Response
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel
from sqlalchemy.orm import Session

from crm.db.modelos import GrupoEconomico, Oportunidade, RegistroDeAlteracao
from crm.domain import indicadores as regras_de_indicadores
from crm.domain.listas import LinhaServico, Situacao
from crm.domain.mrr import mensalizar
from crm.domain.servicos import linha_do_servico

__all__ = ["roteador_de_propostas_em_aberto", "montar"]

ETAPAS = (Situacao.EM_AVALIACAO, Situacao.ON_HOLD, Situacao.ENVIAR_PROPOSTA)
DIAS_DE_ON_HOLD = 60
ZERO = Decimal("0.00")
CENTAVOS = Decimal("0.01")


class PropostaAberta(BaseModel):
    id: int
    grupo: str
    nome: str
    servico: str | None
    etapa: str
    temperatura: str | None
    captador: str | None
    anual: Decimal
    mensalizado: Decimal
    recorrente: bool
    em_on_hold_desde: date | None = None
    on_hold_pela_originacao: bool = False
    dias_em_on_hold: int | None = None


class LinhaDoQuadro(BaseModel):
    etapa: str
    quantas: int
    anual: Decimal
    mensalizado: Decimal
    recorrente: Decimal
    pontual: Decimal
    quente: Decimal
    morno: Decimal
    frio: Decimal
    sem_classificacao: Decimal


class Conversao(BaseModel):
    rotulo: str
    aceitas: int
    decididas: int
    percentual: Decimal | None


class QuadroDePropostas(BaseModel):
    hoje: date
    linhas: list[LinhaDoQuadro]
    total: LinhaDoQuadro
    conversao_do_ano: Conversao
    conversao_12_meses: Conversao
    enviadas_no_mes: int
    enviadas_no_mes_em_aberto: Decimal
    """Valor anual das enviadas no mês que seguem em aberto."""
    mes_das_enviadas: str
    on_hold_antigas: int
    on_hold_antigas_mensalizado: Decimal
    resumo: list[str]
    propostas: list[PropostaAberta]


def _reais(v: Decimal) -> str:
    return "R$ " + f"{v:,.2f}".replace(",", "@").replace(".", ",").replace("@", ".")


def _pct(parte: Decimal, todo: Decimal) -> Decimal | None:
    return (parte / todo * 100).quantize(Decimal("0.1")) if todo > 0 else None


def _recorrente(o: Oportunidade) -> bool:
    linha = o.linha_servico or linha_do_servico(o.servico)
    return linha is not LinhaServico.C2 and o.preco_mensal is not None and o.preco_mensal > 0


def _mensalizado(o: Oportunidade) -> Decimal:
    if _recorrente(o):
        return mensalizar(o.preco_mensal) or ZERO
    return (Decimal(o.preco_anual or 0) / 12).quantize(CENTAVOS)


def _linha(etapa: str, props: list[PropostaAberta]) -> LinhaDoQuadro:
    def soma(f) -> Decimal:
        return sum((p.mensalizado for p in props if f(p)), ZERO)
    return LinhaDoQuadro(
        etapa=etapa, quantas=len(props), anual=sum((p.anual for p in props), ZERO),
        mensalizado=soma(lambda p: True), recorrente=soma(lambda p: p.recorrente), pontual=soma(lambda p: not p.recorrente),
        quente=soma(lambda p: p.temperatura == "Quente"), morno=soma(lambda p: p.temperatura == "Morno"),
        frio=soma(lambda p: p.temperatura == "Frio"), sem_classificacao=soma(lambda p: p.temperatura is None),
    )


def _conversao(rotulo: str, oportunidades: list[Oportunidade]) -> Conversao:
    c = regras_de_indicadores.calcular(oportunidades).taxa_de_conversao
    return Conversao(rotulo=rotulo, aceitas=c.aceitas, decididas=c.decididas, percentual=c.percentual)


def montar(sessao: Session, hoje: date) -> QuadroDePropostas:
    nomes = dict(sessao.execute(sa.select(GrupoEconomico.id, GrupoEconomico.nome)).all())
    todas = list(sessao.scalars(sa.select(Oportunidade).order_by(Oportunidade.id)))
    abertas = [o for o in todas if o.situacao in ETAPAS]

    entrada_em_on_hold: dict[int, date] = {}
    for registro_id, quando in sessao.execute(
        sa.select(RegistroDeAlteracao.registro_id, sa.func.max(RegistroDeAlteracao.quando))
        .where(RegistroDeAlteracao.tabela == "oportunidade", RegistroDeAlteracao.campo == "situacao",
               RegistroDeAlteracao.depois == Situacao.ON_HOLD.value)
        .group_by(RegistroDeAlteracao.registro_id)
    ):
        entrada_em_on_hold[registro_id] = quando.date()

    propostas = []
    for o in abertas:
        p = PropostaAberta(
            id=o.id, grupo=nomes.get(o.grupo_id, "?"), nome=o.nome, servico=o.servico, etapa=o.situacao.value,
            temperatura=o.temperatura.value if o.temperatura else None, captador=o.captador,
            anual=Decimal(o.preco_anual or 0).quantize(CENTAVOS), mensalizado=_mensalizado(o), recorrente=_recorrente(o),
        )
        if o.situacao is Situacao.ON_HOLD:
            desde = entrada_em_on_hold.get(o.id) or o.data_colocacao
            p.em_on_hold_desde = desde
            p.on_hold_pela_originacao = o.id not in entrada_em_on_hold
            p.dias_em_on_hold = (hoje - desde).days if desde else None
        propostas.append(p)

    linhas = [_linha(e.value, [p for p in propostas if p.etapa == e.value]) for e in ETAPAS]
    total = _linha("Total em aberto", propostas)

    inicio_do_ano = date(hoje.year, 1, 1)
    ano = [o for o in todas if o.data_colocacao and inicio_do_ano <= o.data_colocacao <= hoje]
    doze = [o for o in todas if o.data_colocacao and hoje - timedelta(days=365) < o.data_colocacao <= hoje]
    enviadas = [o for o in todas if o.data_envio_proposta and (o.data_envio_proposta.year, o.data_envio_proposta.month)
                == (hoje.year, hoje.month)]
    enviadas_abertas = sum((Decimal(o.preco_anual or 0) for o in enviadas if o.situacao in ETAPAS), ZERO)
    antigas = [p for p in propostas if p.dias_em_on_hold is not None and p.dias_em_on_hold > DIAS_DE_ON_HOLD]
    antigas_mensal = sum((p.mensalizado for p in antigas), ZERO)

    resumo = []
    if total.mensalizado > 0:
        frase = (f"Do mensalizado em aberto ({_reais(total.mensalizado)}/mês), {_reais(total.quente)} está quente "
                 f"({str(_pct(total.quente, total.mensalizado)).replace('.', ',')}%)")
        maior = max(linhas, key=lambda l: l.quente)
        if total.quente > 0 and maior.quente / total.quente >= Decimal("0.5"):
            frase += f", {str(_pct(maior.quente, total.quente)).replace('.', ',')}% dele em \"{maior.etapa}\""
        resumo.append(frase + ".")
    on_hold = next(l for l in linhas if l.etapa == Situacao.ON_HOLD.value)
    resumo.append(f"On hold: {on_hold.quantas} proposta(s), {len(antigas)} há mais de {DIAS_DE_ON_HOLD} dias "
                  f"({_reais(antigas_mensal)}/mês).")
    resumo.append(f"Enviadas no mês: {len(enviadas)}, com {_reais(enviadas_abertas)} (valor anual) ainda em aberto.")

    return QuadroDePropostas(
        hoje=hoje, linhas=linhas, total=total,
        conversao_do_ano=_conversao(f"Conversão · {hoje.year}", ano),
        conversao_12_meses=_conversao("Conversão · últ. 12 meses", doze),
        enviadas_no_mes=len(enviadas), enviadas_no_mes_em_aberto=enviadas_abertas, mes_das_enviadas=f"{hoje:%m/%Y}",
        on_hold_antigas=len(antigas), on_hold_antigas_mensalizado=antigas_mensal, resumo=resumo, propostas=propostas,
    )


def planilha(q: QuadroDePropostas) -> bytes:
    """O mesmo quadro em Excel, para Eduardo ajustar antes de mandar (aba Quadro e aba Propostas)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Quadro"
    ws.append([f"Propostas em aberto · visão gerencial · {q.hoje:%d/%m/%Y}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    cabecalho = ["Etapa", "Qtd", "Valor anual", "Mensalizado", "Quente", "Morno", "Frio", "Sem class."]
    ws.append(cabecalho)
    for c in ws[3]:
        c.font = Font(bold=True)
    for l in [*q.linhas, q.total]:
        ws.append([l.etapa, l.quantas, float(l.anual), float(l.mensalizado), float(l.quente), float(l.morno),
                   float(l.frio), float(l.sem_classificacao)])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for linha in ws.iter_rows(min_row=4, min_col=3, max_col=8):
        for c in linha:
            c.number_format = 'R$ #,##0.00'
    ws.append([])
    ws.append(["Mensalizado: recorrente pela parcela × 13 ÷ 12; pontual pelo valor único ÷ 12."])
    ws.append([])
    for cv in (q.conversao_do_ano, q.conversao_12_meses):
        ws.append([cv.rotulo, f"{cv.percentual}%".replace(".", ",") if cv.percentual is not None else "sem dado",
                   f"{cv.aceitas} aceitas / {cv.decididas} decididas"])
    ws.append([f"Enviadas em {q.mes_das_enviadas}", q.enviadas_no_mes, float(q.enviadas_no_mes_em_aberto)])
    ws.append([f"On hold há mais de {DIAS_DE_ON_HOLD} dias", q.on_hold_antigas, float(q.on_hold_antigas_mensalizado)])
    ws.append([])
    ws.append(["O que o quadro mostra"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for frase in q.resumo:
        ws.append([frase])
    ws.column_dimensions["A"].width = 34
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 16

    wp = wb.create_sheet("Propostas")
    wp.append(["Grupo", "Proposta", "Serviço", "Etapa", "Temperatura", "Captador", "Valor anual", "Mensalizado",
               "Em On hold desde", "Dias em On hold"])
    for c in wp[1]:
        c.font = Font(bold=True)
    for p in q.propostas:
        wp.append([p.grupo, p.nome, p.servico, p.etapa, p.temperatura or "", p.captador or "", float(p.anual),
                   float(p.mensalizado), p.em_on_hold_desde, p.dias_em_on_hold])
    for linha in wp.iter_rows(min_row=2, min_col=7, max_col=8):
        for c in linha:
            c.number_format = 'R$ #,##0.00'
    for linha in wp.iter_rows(min_row=2, min_col=9, max_col=9):
        for c in linha:
            c.number_format = "dd/mm/yyyy"
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def roteador_de_propostas_em_aberto(obter_sessao: Callable[[], Iterator[Session]]) -> APIRouter:
    r = APIRouter(tags=["inteligência de conversão"])

    @r.get("/api/inteligencia/propostas-em-aberto", response_model=QuadroDePropostas)
    def ver(sessao: Session = Depends(obter_sessao, scope="function"), hoje: date | None = None) -> QuadroDePropostas:
        """O quadro gerencial das propostas em aberto. `hoje` existe para teste."""
        return montar(sessao, hoje or date.today())

    @r.get("/api/inteligencia/propostas-em-aberto/planilha")
    def baixar(sessao: Session = Depends(obter_sessao, scope="function"), hoje: date | None = None) -> Response:
        dia = hoje or date.today()
        return Response(
            planilha(montar(sessao, dia)),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="propostas-em-aberto-{dia:%Y%m%d}.xlsx"'},
        )

    return r
