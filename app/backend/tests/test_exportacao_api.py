"""O endpoint que baixa o histórico completo da carteira em Excel (28/09/2026)."""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session, sessionmaker

from crm.api.app import criar_app
from crm.db.modelos import ClassificacaoDoGrupo, Contrato, Empresa, GrupoEconomico
from crm.domain import classificacao as regra
from crm.domain.listas import SituacaoContrato, SituacaoGrupo


@pytest.fixture
def cliente(engine):
    fabrica = sessionmaker(bind=engine, expire_on_commit=False, future=True)
    with TestClient(criar_app(fabrica)) as c:
        yield c


def _snap(sessao: Session, grupo: GrupoEconomico, ref: date, receita: str, n: regra.Notas, revisao: int = 1) -> None:
    pontos = regra.score(n)
    letra = regra.classe(pontos)
    sessao.add(ClassificacaoDoGrupo(
        grupo_id=grupo.id, referencia=ref, revisao=revisao, fonte="teste", versao_dos_parametros=regra.PARAMETROS.versao,
        receita_mensal=Decimal(receita), nota_receita=n.receita, nota_rentabilidade=n.rentabilidade,
        complexidade=n.complexidade, disciplina=n.disciplina, risco_tecnico=n.risco, cross_sell=n.cross_sell,
        adimplencia=n.adimplencia, semaforo=n.semaforo, churn=n.churn, score=pontos, classe=letra,
        classe_efetiva=regra.classe_efetiva(letra, n), alerta_de_churn=regra.alerta_de_churn(letra, n),
        em_cobranca=regra.cobranca(n), eixo_de_acao=regra.eixo_de_acao(letra, n),
    ))


def test_sem_nenhuma_classificacao_devolve_planilha_so_com_cabecalho(cliente):
    r = cliente.get("/api/carteira/exportar")
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "carteira-historico-" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Histórico da Carteira", "Parâmetros", "Critérios e Fórmulas"]
    assert wb["Histórico da Carteira"]["A5"].value is None


def test_traz_todo_o_historico_de_todos_os_grupos(cliente, sessao: Session):
    n = regra.Notas(receita=5, rentabilidade=1, complexidade=3, disciplina=3, risco=3, cross_sell=3, adimplencia=3, semaforo=1, churn=1)
    a = GrupoEconomico(nome="Alfa", situacao=SituacaoGrupo.CLIENTE)
    b = GrupoEconomico(nome="Beta", situacao=SituacaoGrupo.CLIENTE)
    sessao.add_all([a, b]); sessao.flush()
    _snap(sessao, a, date(2026, 6, 30), "1000", n, revisao=1)
    _snap(sessao, a, date(2026, 7, 31), "1000", n, revisao=1)
    _snap(sessao, b, date(2026, 7, 31), "2000", n, revisao=1)
    sessao.commit()

    r = cliente.get("/api/carteira/exportar")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb["Histórico da Carteira"]
    nomes = [ws.cell(row=linha, column=1).value for linha in (5, 6, 7)]
    assert nomes == ["▸ Alfa", "▸ Alfa", "▸ Beta"]  # ordenado por grupo, depois referência — as duas leituras de Alfa antes de Beta
    assert ws["O5"].value.startswith("=F5*")  # fórmula viva do Score


def test_empresas_do_grupo_com_a_mensalidade_dos_contratos_vigentes(cliente, sessao: Session):
    n = regra.Notas(receita=5, rentabilidade=1, complexidade=3, disciplina=3, risco=3, cross_sell=3, adimplencia=3, semaforo=1, churn=1)
    a = GrupoEconomico(nome="Alfa", situacao=SituacaoGrupo.CLIENTE)
    homonimo = GrupoEconomico(nome="Alfa", situacao=SituacaoGrupo.CLIENTE)
    sessao.add_all([a, homonimo]); sessao.flush()
    e1 = Empresa(grupo_id=a.id, razao_social="Alfa Ltda", cnpj="11111111000191")
    e2 = Empresa(grupo_id=a.id, razao_social="Alfa Serviços")
    e3 = Empresa(grupo_id=homonimo.id, razao_social="Outra Alfa")
    sessao.add_all([e1, e2, e3]); sessao.flush()
    sessao.add_all([
        Contrato(grupo_id=a.id, empresa_id=e1.id, anterior_ao_crm=True, situacao=SituacaoContrato.ATIVO, preco_mensal=Decimal("1000.00")),
        Contrato(grupo_id=a.id, empresa_id=e1.id, anterior_ao_crm=True, situacao=SituacaoContrato.SUSPENSO, preco_mensal=Decimal("500.00")),
        Contrato(grupo_id=a.id, empresa_id=e1.id, anterior_ao_crm=True, situacao=SituacaoContrato.ENCERRADO, preco_mensal=Decimal("9999.00")),
        Contrato(grupo_id=a.id, empresa_id=e2.id, anterior_ao_crm=True, situacao=SituacaoContrato.AGUARDANDO_ASSINATURA, preco_mensal=Decimal("300.00")),
    ])
    _snap(sessao, a, date(2026, 7, 31), "1500", n)
    _snap(sessao, homonimo, date(2026, 7, 31), "700", n)
    sessao.commit()

    ws = openpyxl.load_workbook(io.BytesIO(cliente.get("/api/carteira/exportar").content))["Histórico da Carteira"]
    assert [ws.cell(row=r, column=1).value for r in range(5, 10)] == [
        "▸ Alfa", "    • Alfa Ltda", "    • Alfa Serviços", "▸ Alfa", "    • Outra Alfa",
    ]
    assert ws["D6"].value == 1500.0  # ativo + suspenso; encerrado fica de fora
    assert ws["D7"].value is None    # aguardando assinatura ainda não é mensalidade
    assert ws["C6"].value == "11111111000191"


_MIX_POR_PORTE = {
    "Micro": {"Sócio Sênior": "0", "Sócio Júnior/Gerente": "0.01", "Supervisor/Especialista": "0.02",
              "Analista Sênior": "0.47", "Analista Pleno": "0.50", "Analista Júnior": "0"},
    "Pequeno": {"Sócio Sênior": "0", "Sócio Júnior/Gerente": "0.01", "Supervisor/Especialista": "0.02",
                "Analista Sênior": "0.47", "Analista Pleno": "0.50", "Analista Júnior": "0"},
    "Médio": {"Sócio Sênior": "0", "Sócio Júnior/Gerente": "0.01", "Supervisor/Especialista": "0.08",
              "Analista Sênior": "0.30", "Analista Pleno": "0.51", "Analista Júnior": "0.10"},
    "Grande": {"Sócio Sênior": "0.01", "Sócio Júnior/Gerente": "0.01", "Supervisor/Especialista": "0.09",
               "Analista Sênior": "0.30", "Analista Pleno": "0.43", "Analista Júnior": "0.16"},
    "Extra Grande": {"Sócio Sênior": "0.02", "Sócio Júnior/Gerente": "0.05", "Supervisor/Especialista": "0.13",
                      "Analista Sênior": "0.30", "Analista Pleno": "0.30", "Analista Júnior": "0.20"},
}


def test_planilha_usa_os_parametros_gravados_no_banco(cliente, sessao: Session):
    n = regra.Notas(receita=5, rentabilidade=1, complexidade=3, disciplina=3, risco=3, cross_sell=3, adimplencia=3, semaforo=1, churn=1)
    a = GrupoEconomico(nome="Alfa", situacao=SituacaoGrupo.CLIENTE)
    sessao.add(a); sessao.flush()
    _snap(sessao, a, date(2026, 7, 31), "1000", n)
    sessao.commit()

    mix = [
        {"porte": porte, "cargo": cargo, "mix_percentual": valor}
        for porte, cargos in _MIX_POR_PORTE.items()
        for cargo, valor in cargos.items()
    ]
    corpo = {
        "autor": "Eduardo Luiz", "motivo": "teste", "peso_receita": "0.30", "peso_rentabilidade": "0.15",
        "peso_cross_sell": "0.12", "peso_complexidade": "0.12", "peso_disciplina": "0.09", "peso_risco": "0.07",
        "peso_adimplencia": "0.15", "corte_a": "3.95", "corte_b": "3.35", "trava_de_adimplencia": 2, "churn_alto": 4,
        "imposto": "0.11", "teto_de_atrito": "1.5", "atrito_nota_1": "0", "atrito_nota_2": "0.05", "atrito_nota_3": "0.10",
        "atrito_nota_4": "0.30", "atrito_nota_5": "0.50", "corte_margem_2": "0.30", "corte_margem_3": "0.45",
        "corte_margem_4": "0.60", "corte_margem_5": "0.70", "horas_micro": 5, "horas_pequeno": 10, "horas_medio": 16,
        "horas_grande": 40, "horas_extra_grande": 80, "taxa_socio_senior": "93.75", "taxa_socio_junior": "106.25",
        "taxa_supervisor": "75.00", "taxa_analista_senior": "50.00", "taxa_analista_pleno": "31.25",
        "taxa_analista_junior": "18.75", "mix": mix,
    }
    assert cliente.post("/api/carteira/parametros", json=corpo).status_code == 201

    r = cliente.get("/api/carteira/exportar")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb["Parâmetros"]["B8"].value == 0.30  # peso_receita novo, não mais 0.20
