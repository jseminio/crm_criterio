"""Exporta a Grade do funil para Excel (pedido de Karine, 28/09/2026).

As mesmas oito colunas que a Grade mostra, nas mesmas oportunidades que os filtros da tela
selecionam — sem limite de página. Valor sai como número e data como data, para a planilha
somar e ordenar sozinha. Nada aqui grava nada; é só leitura formatada.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

__all__ = ["COLUNAS", "LinhaDoFunil", "gerar_planilha_do_funil"]

COLUNAS = ["Cliente", "Oportunidade", "Situação", "Temperatura", "Captador", "Originação", "Mensal", "Anual"]
_LARGURAS = [36, 40, 18, 14, 20, 13, 15, 15]
_CINZA_CLARO = "F2F2F2"
_DINHEIRO = "R$ #,##0.00"


@dataclass(frozen=True)
class LinhaDoFunil:
    """Uma linha da Grade, já com o nome do cliente resolvido."""

    cliente: str | None
    oportunidade: str
    situacao: str
    temperatura: str | None
    captador: str | None
    originacao: date | None
    mensal: Decimal | None
    anual: Decimal | None


def gerar_planilha_do_funil(linhas: list[LinhaDoFunil]) -> bytes:
    """Devolve os bytes do `.xlsx`: uma aba, cabeçalho congelado e filtro ligado."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Funil"

    for col, titulo in enumerate(COLUNAS, start=1):
        c = ws.cell(row=1, column=col, value=titulo)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=_CINZA_CLARO)
        ws.column_dimensions[c.column_letter].width = _LARGURAS[col - 1]

    for linha, l in enumerate(linhas, start=2):
        ws.cell(row=linha, column=1, value=l.cliente)
        ws.cell(row=linha, column=2, value=l.oportunidade)
        ws.cell(row=linha, column=3, value=l.situacao)
        ws.cell(row=linha, column=4, value=l.temperatura)
        ws.cell(row=linha, column=5, value=l.captador)
        ws.cell(row=linha, column=6, value=l.originacao).number_format = "DD/MM/YYYY"
        for col, valor in ((7, l.mensal), (8, l.anual)):
            c = ws.cell(row=linha, column=col, value=float(valor) if valor is not None else None)
            c.number_format = _DINHEIRO

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{max(len(linhas) + 1, 1)}"

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()
