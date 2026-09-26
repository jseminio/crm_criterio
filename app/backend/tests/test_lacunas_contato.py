"""Importação da planilha de lacunas de contato: várias empresas por cliente, sem sobrescrever em silêncio."""

from __future__ import annotations

import pytest
import sqlalchemy as sa
from openpyxl import Workbook
from sqlalchemy.orm import Session

from crm.carga.lacunas_contato import COLUNAS, Relatorio, aplicar, cnpj_valido, ler
from crm.db.modelos import Empresa, GrupoEconomico, PessoaContato

CNPJ_A = "11222333000181"  # válido
CNPJ_B = "11444777000161"  # válido


def _planilha(tmp_path, linhas: list[dict]):
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    ws.append(list(COLUNAS))
    for l in linhas:
        ws.append([l.get(t, "") for t in COLUNAS])
    ws.append(["", "Lacunas restantes"])  # rodapé de totais: deve ser ignorado
    caminho = tmp_path / "p.xlsx"
    wb.save(caminho)
    return caminho


@pytest.fixture
def grupo(sessao: Session) -> GrupoEconomico:
    g = GrupoEconomico(nome="Grupo Alfa")
    sessao.add(g)
    sessao.commit()
    return g


def _rodar(sessao, caminho, **kw):
    rel = Relatorio()
    linhas = ler(caminho, rel)
    aplicar(sessao, linhas, rel, **kw)
    return rel


def test_cnpj_valido():
    assert cnpj_valido(CNPJ_A) and cnpj_valido(CNPJ_B)
    assert not cnpj_valido("11222333000182") and not cnpj_valido("11111111111111") and not cnpj_valido("123")


def test_varias_empresas_do_mesmo_cliente_em_linhas_separadas(sessao, grupo, tmp_path):
    p = _planilha(tmp_path, [
        {"ID_GRUPO": grupo.id, "Razão social": "Alfa Comércio Ltda", "CNPJ (só números)": CNPJ_A,
         "Logradouro": "Rua Ação", "Número / compl.": "10 sala 2", "Bairro": "Centro",
         "Município": "Rio de Janeiro", "UF": "rj", "CEP": "20040-020",
         "Contato — nome": "Maria", "E-mail": "MARIA@alfa.com", "Telefone": "21 99999-0000"},
        {"ID_GRUPO": grupo.id, "Razão social": "Alfa Serviços SA", "CNPJ (só números)": CNPJ_B},
    ])
    rel = _rodar(sessao, p)
    assert rel.pode_aplicar and rel.empresas_criadas == 2 and rel.contatos_criados == 1
    e = sessao.scalars(sa.select(Empresa).where(Empresa.cnpj == CNPJ_A)).one()
    assert (e.numero, e.complemento, e.uf, e.cep) == ("10", "sala 2", "RJ", "20040020")
    c = sessao.scalars(sa.select(PessoaContato)).one()
    assert c.empresa_id == e.id and c.email == "maria@alfa.com"


def test_rodar_duas_vezes_nao_duplica(sessao, grupo, tmp_path):
    p = _planilha(tmp_path, [{"ID_GRUPO": grupo.id, "Razão social": "Alfa Ltda", "CNPJ (só números)": CNPJ_A,
                              "Contato — nome": "Maria", "E-mail": "m@alfa.com"}])
    _rodar(sessao, p)
    rel = _rodar(sessao, p)
    assert rel.empresas_criadas == 0 and rel.contatos_criados == 0
    assert sessao.scalar(sa.select(sa.func.count()).select_from(Empresa)) == 1
    assert sessao.scalar(sa.select(sa.func.count()).select_from(PessoaContato)) == 1


def test_valor_diferente_vira_conflito_e_nao_e_sobrescrito(sessao, grupo, tmp_path):
    sessao.add(Empresa(grupo_id=grupo.id, razao_social="Alfa Ltda", cnpj=CNPJ_A, municipio="Niterói"))
    sessao.commit()
    p = _planilha(tmp_path, [{"ID_GRUPO": grupo.id, "CNPJ (só números)": CNPJ_A, "Município": "Rio de Janeiro"}])
    rel = _rodar(sessao, p)
    assert len(rel.conflitos) == 1
    assert sessao.scalars(sa.select(Empresa)).one().municipio == "Niterói"
    _rodar(sessao, p, sobrescrever=True)
    assert sessao.scalars(sa.select(Empresa)).one().municipio == "Rio de Janeiro"


def test_erros_bloqueiam(sessao, grupo, tmp_path):
    p = _planilha(tmp_path, [
        {"ID_GRUPO": grupo.id, "CNPJ (só números)": "123", "UF": "XX", "CEP": "1", "E-mail": "sem-arroba"},
        {"ID_GRUPO": 99999, "Razão social": "Fantasma"},
    ])
    rel = _rodar(sessao, p)
    assert not rel.pode_aplicar
    texto = " ".join(rel.erros)
    for pedaco in ("CNPJ inválido", "UF inválida", "CEP", "e-mail inválido", "não existe grupo"):
        assert pedaco in texto


def test_cnpj_repetido_em_outro_cliente_e_erro(sessao, grupo, tmp_path):
    outro = GrupoEconomico(nome="Beta")
    sessao.add(outro)
    sessao.commit()
    p = _planilha(tmp_path, [{"ID_GRUPO": grupo.id, "CNPJ (só números)": CNPJ_A},
                             {"ID_GRUPO": outro.id, "CNPJ (só números)": CNPJ_A}])
    assert any("outro cliente" in e for e in _rodar(sessao, p).erros)


def test_linha_sem_dado_e_ignorada_e_planilha_sem_colunas_e_recusada(sessao, grupo, tmp_path):
    rel = _rodar(sessao, _planilha(tmp_path, [{"ID_GRUPO": grupo.id}]))
    assert rel.linhas_vazias == 1 and rel.empresas_criadas == 0
    wb = Workbook(); wb.active.title = "Clientes"; wb.active.append(["outra", "coisa"])
    wb.save(tmp_path / "q.xlsx")
    assert not _rodar(sessao, tmp_path / "q.xlsx").pode_aplicar


# ---------------------------------------------------------------- planilha por empresa (26/09/2026)

def _planilha_por_empresa(tmp_path, clientes=(), prospects=()):
    """Aba Clientes com ID_EMPRESA (uma linha por empresa) e aba Prospects (uma por grupo)."""
    base = list(COLUNAS)
    wb = Workbook()
    c = wb.active
    c.title = "Clientes"
    c.append(["ID_EMPRESA"] + base)
    for l in clientes:
        c.append([l.get("ID_EMPRESA", "")] + [l.get(t, "") for t in base])
    p = wb.create_sheet("Prospects")
    p.append(base)
    for l in prospects:
        p.append([l.get(t, "") for t in base])
    caminho = tmp_path / "por_empresa.xlsx"
    wb.save(caminho)
    return caminho


def test_linha_por_empresa_usa_o_id_e_preenche_so_o_que_falta(sessao, grupo, tmp_path):
    e1 = Empresa(grupo_id=grupo.id, razao_social="Alfa Comércio Ltda", cnpj=CNPJ_A)
    e2 = Empresa(grupo_id=grupo.id, razao_social="Alfa Serviços SA", cnpj=CNPJ_B)
    sessao.add_all([e1, e2]); sessao.commit()
    p = _planilha_por_empresa(tmp_path, clientes=[
        {"ID_EMPRESA": e1.id, "ID_GRUPO": grupo.id, "Razão social": "Alfa Comércio Ltda", "CNPJ (só números)": CNPJ_A,
         "Logradouro": "Rua Ação", "Município": "Rio de Janeiro", "UF": "RJ", "CEP": "20040-020",
         "Contato — nome": "Maria", "E-mail": "maria@alfa.com"},
        {"ID_EMPRESA": e2.id, "ID_GRUPO": grupo.id, "Razão social": "Alfa Serviços SA", "CNPJ (só números)": CNPJ_B,
         "Contato — nome": "João"},
    ])
    rel = _rodar(sessao, p)
    assert rel.pode_aplicar and rel.empresas_atualizadas == 1 and rel.contatos_criados == 2
    assert sessao.get(Empresa, e1.id).logradouro == "Rua Ação" and sessao.get(Empresa, e2.id).logradouro is None
    contatos = {c.nome: c.empresa_id for c in sessao.scalars(sa.select(PessoaContato))}
    assert contatos == {"Maria": e1.id, "João": e2.id}


def test_planilha_intacta_nao_muda_nada(sessao, grupo, tmp_path):
    e = Empresa(grupo_id=grupo.id, razao_social="Alfa Ltda", cnpj=CNPJ_A)
    sessao.add(e); sessao.commit()
    p = _planilha_por_empresa(tmp_path, clientes=[{"ID_EMPRESA": e.id, "ID_GRUPO": grupo.id, "Razão social": "Alfa Ltda", "CNPJ (só números)": CNPJ_A}])
    rel = _rodar(sessao, p)
    assert (rel.empresas_atualizadas, rel.contatos_criados, len(rel.conflitos)) == (0, 0, 0)


def test_id_de_empresa_de_outro_grupo_e_erro(sessao, grupo, tmp_path):
    outro = GrupoEconomico(nome="Beta"); sessao.add(outro); sessao.flush()
    e = Empresa(grupo_id=outro.id, razao_social="Beta Ltda", cnpj=CNPJ_B)
    sessao.add(e); sessao.commit()
    rel = _rodar(sessao, _planilha_por_empresa(tmp_path, clientes=[{"ID_EMPRESA": e.id, "ID_GRUPO": grupo.id, "E-mail": "x@y.com"}]))
    assert not rel.pode_aplicar and "não é do grupo" in " ".join(rel.erros)


def test_prospects_e_clientes_sao_lidos_e_o_aviso_diz_a_aba(sessao, grupo, tmp_path):
    p = _planilha_por_empresa(tmp_path, prospects=[{"ID_GRUPO": grupo.id, "Contato — nome": "Bia", "E-mail": "sem-arroba"}])
    rel = _rodar(sessao, p)
    assert any(x.startswith("Prospects, linha 2") for x in rel.erros)


def test_planilha_sem_nenhuma_das_duas_abas_e_recusada(sessao, tmp_path):
    wb = Workbook(); wb.active.title = "Outra"; wb.save(tmp_path / "x.xlsx")
    rel = Relatorio()
    assert ler(tmp_path / "x.xlsx", rel) == [] and not rel.pode_aplicar
