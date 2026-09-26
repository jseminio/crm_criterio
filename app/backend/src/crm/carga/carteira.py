"""Carrega a carteira que já existia antes do CRM, a partir da planilha de saúde da carteira.

Fonte: aba **"4. Clientes"** de `Rentabilidade_Grupo_COMPLETO.xlsx` — uma linha por empresa
(CNPJ), com razão social, grupo econômico (ou "Sem grupo"), escopo contratado e honorário
mensal. **O arquivo fica fora do repositório** (dado de cliente); só o código mora aqui.

Regras (decisões de Eduardo, 26/09/2026):

- Cada linha vira uma **Empresa** (por CNPJ) e um **Contrato Ativo** dela, com o honorário
  como preço mensal. O contrato é marcado `anterior_ao_crm`: **a planilha não traz a data de
  assinatura e não se inventa uma** — assim ele conta no MRR desde sempre, nunca como "novo".
- Empresas do mesmo "Grupo Econômico" ficam no mesmo **grupo**; quem está como "Sem grupo" é
  um grupo de uma empresa só, com o nome da própria razão social.
- Grupo que **já existe** no CRM (vindo das propostas) é reaproveitado quando a *chave do
  nome* é a mesma (`crm.domain.sugestoes_de_fusao.chave_do_nome`); se há mais de um, usa o de
  mais propostas e **avisa**. Nomes diferentes viram grupo novo — a fusão continua sendo um
  clique humano depois. Um grupo Prospect que recebe cliente vira **Cliente**.
- **Idempotente**: rodar de novo não duplica (empresa por CNPJ, contrato por empresa).
- **Nada é sobrescrito**: se o CRM já tem preço diferente, vira conflito no relatório.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

import sqlalchemy as sa
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from crm.carga.lacunas_contato import cnpj_valido
from crm.db.modelos import Contrato, Empresa, GrupoEconomico, Oportunidade
from crm.domain.listas import Origem, SituacaoContrato, SituacaoGrupo
from crm.domain.sugestoes_de_fusao import chave_do_nome

__all__ = ["Linha", "Relatorio", "ler", "aplicar"]

ABA = "4. Clientes"
SEM_GRUPO = "sem grupo"
ESCOPO_A_VERIFICAR = "verificar contrato"
COLUNAS = {
    "Razão Social": "razao_social",
    "CNPJ": "cnpj",
    "Grupo Econômico": "grupo",
    "Escopo Contratado (referência)": "escopo",
    "Honorário (R$/mês)": "honorario",
}


@dataclass(frozen=True)
class Linha:
    n: int
    razao_social: str
    cnpj: str
    grupo: str | None
    """`None` = cliente individual ("Sem grupo")."""
    escopo: str | None
    honorario: Decimal


@dataclass
class Relatorio:
    linhas: int = 0
    total_mensal: Decimal = Decimal("0.00")
    grupos_reaproveitados: int = 0
    grupos_novos: int = 0
    empresas_criadas: int = 0
    empresas_existentes: int = 0
    contratos_criados: int = 0
    contratos_existentes: int = 0
    prospects_que_viraram_cliente: int = 0
    erros: list[str] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)
    conflitos: list[str] = field(default_factory=list)

    @property
    def pode_aplicar(self) -> bool:
        return not self.erros


def _texto(v) -> str:
    return " ".join(str(v).split()) if v is not None else ""


def ler(caminho: Path, rel: Relatorio) -> list[Linha]:
    wb = load_workbook(caminho, data_only=True)
    if ABA not in wb.sheetnames:
        rel.erros.append(f"A planilha não tem a aba “{ABA}”. Use Rentabilidade_Grupo_COMPLETO.xlsx.")
        return []
    ws = wb[ABA]
    # o cabeçalho pode ter quebra de linha dentro da célula
    cabecalho = None
    for r in range(1, 12):
        nomes = {_texto(c.value): i for i, c in enumerate(ws[r]) if c.value is not None}
        if "Razão Social" in nomes and "CNPJ" in nomes:
            cabecalho, indice = r, nomes
            break
    if cabecalho is None:
        rel.erros.append(f"Não achei o cabeçalho (Razão Social, CNPJ…) na aba “{ABA}”.")
        return []
    faltam = [c for c in COLUNAS if c not in indice]
    if faltam:
        rel.erros.append(f"Faltam colunas na aba: {', '.join(faltam)}.")
        return []

    linhas: list[Linha] = []
    vistos: dict[str, int] = {}
    for n, row in enumerate(ws.iter_rows(min_row=cabecalho + 1, values_only=True), cabecalho + 1):
        razao = _texto(row[indice["Razão Social"]])
        if not razao:
            continue
        rel.linhas += 1
        cnpj = re.sub(r"\D", "", _texto(row[indice["CNPJ"]]))
        if not cnpj_valido(cnpj):
            rel.erros.append(f"Linha {n}: CNPJ inválido ({cnpj!r}).")
            continue
        if cnpj in vistos:
            rel.erros.append(f"Linha {n}: o CNPJ {cnpj} já apareceu na linha {vistos[cnpj]}.")
            continue
        vistos[cnpj] = n
        bruto = row[indice["Honorário (R$/mês)"]]
        try:
            honorario = Decimal(str(bruto)).quantize(Decimal("0.01"))
        except Exception:
            rel.erros.append(f"Linha {n}: honorário ilegível ({bruto!r}).")
            continue
        if honorario <= 0:
            rel.erros.append(f"Linha {n}: honorário precisa ser maior que zero.")
            continue
        grupo = _texto(row[indice["Grupo Econômico"]])
        escopo = _texto(row[indice["Escopo Contratado (referência)"]]) or None
        if escopo and ESCOPO_A_VERIFICAR in escopo.lower():
            rel.avisos.append(f"Linha {n}: escopo marcado para verificar no contrato; ficou em branco.")
            escopo = None
        if escopo and len(escopo) > 200:
            rel.erros.append(f"Linha {n}: escopo passa de 200 caracteres.")
            continue
        rel.total_mensal += honorario
        linhas.append(Linha(n, razao, cnpj, None if not grupo or grupo.lower() == SEM_GRUPO else grupo, escopo, honorario))
    return linhas


def _grupo_existente(sessao: Session, nome: str, rel: Relatorio) -> GrupoEconomico | None:
    chave = chave_do_nome(nome)
    if not chave:
        return None
    candidatos = [
        g for g in sessao.scalars(sa.select(GrupoEconomico).where(GrupoEconomico.fundido_em_id.is_(None)))
        if chave_do_nome(g.nome) == chave
    ]
    if not candidatos:
        return None
    if len(candidatos) > 1:
        def propostas(g):
            return sessao.scalar(sa.select(sa.func.count(Oportunidade.id)).where(Oportunidade.grupo_id == g.id))
        candidatos.sort(key=lambda g: (-propostas(g), g.id))
        rel.avisos.append(
            f"“{nome}” combina com {len(candidatos)} grupos do CRM; usei o de mais propostas "
            f"(id {candidatos[0].id}). Junte os demais depois, em Grupos."
        )
    return candidatos[0]


def aplicar(sessao: Session, linhas: list[Linha], rel: Relatorio) -> None:
    """Grava. Chamar dentro de uma transação que o chamador confirma ou desfaz."""
    grupos: dict[str, GrupoEconomico] = {}  # cache por nome da carteira
    for lin in linhas:
        nome_do_grupo = lin.grupo or lin.razao_social
        grupo = grupos.get(nome_do_grupo)
        if grupo is None:
            grupo = _grupo_existente(sessao, nome_do_grupo, rel)
            if grupo is None:
                grupo = GrupoEconomico(nome=nome_do_grupo, situacao=SituacaoGrupo.CLIENTE, origem=Origem.CARTEIRA_ANTERIOR)
                sessao.add(grupo)
                sessao.flush()
                rel.grupos_novos += 1
            else:
                rel.grupos_reaproveitados += 1
                if grupo.situacao is SituacaoGrupo.PROSPECT:
                    grupo.situacao = SituacaoGrupo.CLIENTE
                    rel.prospects_que_viraram_cliente += 1
            grupos[nome_do_grupo] = grupo

        empresa = sessao.scalars(sa.select(Empresa).where(Empresa.cnpj == lin.cnpj)).first()
        if empresa is None:
            empresa = Empresa(grupo_id=grupo.id, razao_social=lin.razao_social, cnpj=lin.cnpj)
            sessao.add(empresa)
            sessao.flush()
            rel.empresas_criadas += 1
        else:
            rel.empresas_existentes += 1
            if empresa.grupo_id != grupo.id:
                rel.conflitos.append(f"Linha {lin.n}: o CNPJ {lin.cnpj} já está no grupo {empresa.grupo_id}, e a planilha o põe em “{nome_do_grupo}” (não movido).")

        contrato = sessao.scalars(sa.select(Contrato).where(Contrato.empresa_id == empresa.id)).first()
        if contrato is None:
            sessao.add(
                Contrato(
                    grupo_id=empresa.grupo_id, empresa_id=empresa.id, anterior_ao_crm=True,
                    escopo=lin.escopo, preco_mensal=lin.honorario, situacao=SituacaoContrato.ATIVO,
                )
            )
            rel.contratos_criados += 1
        else:
            rel.contratos_existentes += 1
            if contrato.preco_mensal != lin.honorario:
                rel.conflitos.append(f"Linha {lin.n}: o contrato já tem preço {contrato.preco_mensal} e a planilha traz {lin.honorario} (não alterado).")
    sessao.flush()
