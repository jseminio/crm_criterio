"""Exporta o histórico completo da classificação da carteira para Excel, para conferência
(pedido de Eduardo, 28/09/2026): "planilha em excel com fórmulas auditáveis".

Segue o mesmo modelo de `Faturamento_Grupo_COMPLETO.xlsx`: uma aba com a tabela e uma fórmula viva
para o que é simples de auditar (aqui, Score e Classe — não só a nota de receita como no modelo
de faturamento, porque o Score inteiro é uma soma ponderada direta), uma aba "Parâmetros" com os
pesos e cortes que a fórmula referencia, e uma aba "Critérios e Fórmulas" documentando o resto
(classe efetiva, alerta de churn, eixo de ação) que **não** vira fórmula — são regras com texto e
precedência, não conta; ficam como valor, iguais ao que a tela mostra.

**Todo o histórico, não só o snapshot atual**: cada linha de `ClassificacaoDoGrupo` que já existiu,
pela mesma razão do banco ser imutável — permite comparar mês a mês. Nada aqui grava nada; é só
leitura formatada.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

__all__ = ["LinhaDeHistorico", "ParametrosDaPlanilha", "gerar_planilha"]

_AZUL_NOTURNO = "0B0E17"
_BRANCO = "FFFFFF"
_CINZA_CLARO = "F2F2F2"


@dataclass(frozen=True)
class LinhaDeHistorico:
    """Uma leitura de `ClassificacaoDoGrupo`, com o nome do grupo já resolvido."""

    grupo_nome: str
    referencia: date
    revisao: int
    registrado_em: datetime
    fonte: str
    atribuido_por: str | None
    motivo: str | None
    receita_mensal: Decimal
    nota_receita: Decimal
    nota_rentabilidade: Decimal
    complexidade: Decimal
    disciplina: Decimal
    risco_tecnico: Decimal
    cross_sell: Decimal
    adimplencia: Decimal
    semaforo: int
    churn: int | None
    score: Decimal
    classe: str
    classe_efetiva: str
    alerta_de_churn: str | None
    em_cobranca: bool
    eixo_de_acao: str


@dataclass(frozen=True)
class ParametrosDaPlanilha:
    """Os campos de `VersaoDeParametros` que a aba "Parâmetros" mostra e a fórmula do Score usa."""

    corte_a: Decimal
    corte_b: Decimal
    trava_de_adimplencia: int
    churn_alto: int
    peso_receita: Decimal
    peso_rentabilidade: Decimal
    peso_adimplencia: Decimal
    peso_complexidade: Decimal
    peso_cross_sell: Decimal
    peso_disciplina: Decimal
    peso_risco: Decimal


_COLUNAS = (
    "Grupo", "Referência", "Revisão", "Registrado em", "Fonte", "Atribuído por", "Motivo",
    "Receita mensal (R$)", "Nota Receita", "Nota Rentabilidade", "Complexidade", "Disciplina",
    "Risco técnico", "Cross-sell", "Adimplência", "Semáforo", "Churn", "Score", "Classe",
    "Classe efetiva", "Alerta de churn", "Em cobrança", "Eixo de ação",
)
_LINHA_DO_CABECALHO = 4
_PRIMEIRA_LINHA_DE_DADOS = 5


def _aba_de_parametros(wb: Workbook, p: ParametrosDaPlanilha) -> None:
    ws = wb.create_sheet("Parâmetros")
    ws["A1"] = "PARÂMETROS DE CLASSIFICAÇÃO"
    ws["A1"].font = Font(bold=True, size=13)
    linhas = (
        (2, "Corte A — Score mínimo", float(p.corte_a)),
        (3, "Corte B — Score mínimo", float(p.corte_b)),
        (4, "Trava inadimplência ≤", p.trava_de_adimplencia),
        (5, "Churn alto ≥", p.churn_alto),
    )
    for linha, rotulo, valor in linhas:
        ws.cell(row=linha, column=1, value=rotulo)
        ws.cell(row=linha, column=2, value=valor)
    ws["A7"] = "Pesos do Score"
    ws["A7"].font = Font(bold=True)
    pesos = (
        (8, "Receita", p.peso_receita), (9, "Rentabilidade", p.peso_rentabilidade),
        (10, "Adimplência", p.peso_adimplencia), (11, "Complexidade (invertido)", p.peso_complexidade),
        (12, "Cross-sell", p.peso_cross_sell), (13, "Disciplina", p.peso_disciplina),
        (14, "Risco técnico (invertido)", p.peso_risco),
    )
    for linha, rotulo, valor in pesos:
        ws.cell(row=linha, column=1, value=f"  {rotulo}")
        celula = ws.cell(row=linha, column=2, value=float(valor))
        celula.number_format = "0%"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 14


def _aba_de_criterios(wb: Workbook) -> None:
    ws = wb.create_sheet("Critérios e Fórmulas")
    ws["B1"] = "CRITÉRIOS E FÓRMULAS — AUDITORIA DO HISTÓRICO DA CARTEIRA"
    ws["B1"].font = Font(bold=True, size=12)
    linhas = (
        ("Score", "FÓRMULA VIVA: soma ponderada das sete notas (Complexidade e Risco técnico entram invertidos, "
                   "6 − nota). Pesos e cortes na aba Parâmetros — mude lá e o Score recalcula sozinho."),
        ("Classe", "FÓRMULA VIVA: A se Score ≥ Corte A | B se Score ≥ Corte B | senão C."),
        ("Classe efetiva", "Letra da Classe + semáforo (ex. \"B3\"); com \"(TRAVADO)\" se Adimplência ≤ trava. "
                            "Não é fórmula aqui: mistura texto e a mesma regra de trava do Eixo de Ação."),
        ("Alerta de churn", "\"⚠\" se churn alto em Classe A/B, \"⚑\" se em Classe C, vazio se não. Valor, não fórmula."),
        ("Eixo de ação", "Cobrança > alerta de churn > o resto, nessa ordem de precedência. Valor, não fórmula "
                          "(regra de texto com precedência, não conta)."),
        ("Histórico", "Uma linha por leitura de `ClassificacaoDoGrupo` — o banco é imutável, uma edição de nota "
                       "grava linha nova e a anterior fica. Esta planilha traz **todas**, não só a mais recente."),
    )
    for i, (rotulo, texto) in enumerate(linhas):
        linha = 3 + i
        ws.cell(row=linha, column=2, value=rotulo).font = Font(bold=True)
        ws.cell(row=linha, column=3, value=texto).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 90


def _cabecalho(ws: Worksheet) -> None:
    ws["A1"] = "HISTÓRICO DA CLASSIFICAÇÃO DA CARTEIRA — todas as leituras, não só a mais recente"
    ws["A1"].font = Font(bold=True, size=13, color=_BRANCO)
    ws["A2"] = "Score e Classe têm fórmula viva (referencia a aba Parâmetros); o resto é o valor já calculado no CRM."
    ws["A2"].font = Font(italic=True, color=_BRANCO)
    for linha in (1, 2):
        for col in range(1, len(_COLUNAS) + 1):
            ws.cell(row=linha, column=col).fill = PatternFill("solid", fgColor=_AZUL_NOTURNO)
    for col, titulo in enumerate(_COLUNAS, start=1):
        c = ws.cell(row=_LINHA_DO_CABECALHO, column=col, value=titulo)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=_CINZA_CLARO)


def gerar_planilha(linhas: list[LinhaDeHistorico], parametros: ParametrosDaPlanilha) -> bytes:
    """Devolve os bytes do `.xlsx` — três abas, prontas para `Content-Disposition: attachment`."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Histórico da Carteira"
    _cabecalho(ws)

    for i, l in enumerate(linhas):
        linha = _PRIMEIRA_LINHA_DE_DADOS + i
        ws.cell(row=linha, column=1, value=l.grupo_nome)
        ws.cell(row=linha, column=2, value=l.referencia).number_format = "DD/MM/AAAA"
        ws.cell(row=linha, column=3, value=l.revisao)
        c = ws.cell(row=linha, column=4, value=l.registrado_em.replace(tzinfo=None))
        c.number_format = "DD/MM/AAAA HH:MM"
        ws.cell(row=linha, column=5, value=l.fonte)
        ws.cell(row=linha, column=6, value=l.atribuido_por or "")
        ws.cell(row=linha, column=7, value=l.motivo or "")
        ws.cell(row=linha, column=8, value=float(l.receita_mensal)).number_format = "R$ #,##0.00"
        for col, valor in (
            (9, l.nota_receita), (10, l.nota_rentabilidade), (11, l.complexidade), (12, l.disciplina),
            (13, l.risco_tecnico), (14, l.cross_sell), (15, l.adimplencia),
        ):
            ws.cell(row=linha, column=col, value=float(valor)).number_format = "0.00"
        ws.cell(row=linha, column=16, value=l.semaforo)
        ws.cell(row=linha, column=17, value=l.churn if l.churn is not None else "")

        # Score: soma ponderada, referenciando os pesos da aba Parâmetros — mesma fórmula de
        # `crm.domain.classificacao.score`. Complexidade (K) e Risco técnico (M) entram invertidos.
        formula_do_score = (
            f"=I{linha}*Parâmetros!$B$8 + J{linha}*Parâmetros!$B$9 + O{linha}*Parâmetros!$B$10"
            f" + (6-K{linha})*Parâmetros!$B$11 + N{linha}*Parâmetros!$B$12 + L{linha}*Parâmetros!$B$13"
            f" + (6-M{linha})*Parâmetros!$B$14"
        )
        ws.cell(row=linha, column=18, value=formula_do_score).number_format = "0.0000"
        # Classe: os dois cortes fixos da aba Parâmetros — mesma fórmula de `classificacao.classe`.
        formula_da_classe = f'=IF(R{linha}>=Parâmetros!$B$2,"A",IF(R{linha}>=Parâmetros!$B$3,"B","C"))'
        ws.cell(row=linha, column=19, value=formula_da_classe)

        ws.cell(row=linha, column=20, value=l.classe_efetiva)
        ws.cell(row=linha, column=21, value=l.alerta_de_churn or "")
        ws.cell(row=linha, column=22, value="Sim" if l.em_cobranca else "Não")
        ws.cell(row=linha, column=23, value=l.eixo_de_acao)

    larguras = (26, 12, 9, 16, 16, 16, 30, 16, 11, 13, 12, 11, 12, 10, 11, 10, 8, 9, 8, 13, 12, 11, 32)
    for col, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(col)].width = largura
    ws.freeze_panes = f"A{_PRIMEIRA_LINHA_DE_DADOS}"

    _aba_de_parametros(wb, parametros)
    _aba_de_criterios(wb)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
