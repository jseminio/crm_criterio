"""MRR da carteira por cliente, contratado × recebido e a importação da planilha de recebimentos
(10/10/2026). A lista que a tela abre precisa somar exatamente o número do card."""

from __future__ import annotations

import base64
from datetime import date
from decimal import Decimal as D

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
import sqlalchemy as sa
from sqlalchemy.orm import Session, sessionmaker

from crm.acesso.entrada import ConfiguracaoDeEntrada
from crm.api.app import criar_app
from crm.db.modelos import Contrato, Empresa, EventoDeContrato, GrupoEconomico, MetaDeIndicador, Perfil, Recebimento, Usuario
from crm.domain import recebimentos as regra
from crm.domain.listas import IniciativaDoEncerramento, SituacaoContrato, TipoDeEventoDeContrato as T
from crm.manutencao.tarefas import meta_de_mrr_em_13_parcelas

CONFIG = ConfiguracaoDeEntrada("t", "c", frozenset({"eduardo@grupocriterio.com.br"}))
ADMIN = {"Authorization": "Bearer eduardo@grupocriterio.com.br|Eduardo Luiz"}
KARINE = {"Authorization": "Bearer karine@grupocriterio.com.br|Karine N"}
HOJE = "2026-10-20"


def validar(token: str) -> dict:
    email, nome = token.split("|", 1)
    return {"preferred_username": email, "name": nome}


@pytest.fixture
def base(engine):
    with Session(engine) as s:
        s.add(Perfil(nome="Administrador", administrador=True, permissoes=[]))
        comercial = Perfil(nome="Comercial", permissoes=["funil.ver", "contratos.ver", "configuracoes.metas"])
        alfa, beta, gama = GrupoEconomico(nome="Grupo Alfa"), GrupoEconomico(nome="Beta"), GrupoEconomico(nome="Gama")
        s.add_all([comercial, alfa, beta, gama])
        s.flush()
        s.add(Usuario(email="karine@grupocriterio.com.br", perfil_id=comercial.id))
        ea = Empresa(grupo_id=alfa.id, razao_social="Alfa Comércio Ltda", cnpj="11222333000181")
        eb = Empresa(grupo_id=beta.id, razao_social="Beta Serviços")
        s.add_all([ea, eb])
        s.flush()
        a1 = Contrato(grupo_id=alfa.id, empresa_id=ea.id, anterior_ao_crm=True, situacao=SituacaoContrato.ATIVO,
                      preco_mensal=D("12000"), escopo="Contábil")
        a2 = Contrato(grupo_id=alfa.id, situacao=SituacaoContrato.ATIVO, preco_mensal=D("2400"),
                      data_inicio=date(2026, 10, 5), escopo="BPO Financeiro")
        b1 = Contrato(grupo_id=beta.id, empresa_id=eb.id, anterior_ao_crm=True, situacao=SituacaoContrato.ATIVO,
                      preco_mensal=D("3600"), escopo="Contábil")
        g1 = Contrato(grupo_id=gama.id, situacao=SituacaoContrato.ENCERRADO, preco_mensal=D("6000"),
                      data_inicio=date(2025, 6, 1), escopo="Contábil")
        susp = Contrato(grupo_id=beta.id, situacao=SituacaoContrato.SUSPENSO, preco_mensal=D("1200"),
                        data_inicio=date(2026, 1, 1), escopo="DP")
        s.add_all([a1, a2, b1, g1, susp])
        s.flush()
        s.add(EventoDeContrato(contrato_id=b1.id, tipo=T.REAJUSTE, data_do_evento=date(2026, 10, 10),
                               preco_mensal_anterior=D("3000"), preco_mensal_novo=D("3600")))
        s.add(EventoDeContrato(contrato_id=g1.id, tipo=T.ENCERRAMENTO, data_do_evento=date(2026, 10, 15),
                               iniciativa=IniciativaDoEncerramento.CLIENTE))
        s.commit()


@pytest.fixture
def cliente(engine, base):
    fabrica = sessionmaker(bind=engine, expire_on_commit=False, future=True)
    with TestClient(criar_app(fabrica, entrada=CONFIG, validar_token=validar)) as c:
        yield c


def _xlsx(linhas: list[list[object]]) -> str:
    livro = Workbook()
    folha = livro.active
    for l in linhas:
        folha.append(l)
    import io

    saida = io.BytesIO()
    livro.save(saida)
    return base64.b64encode(saida.getvalue()).decode()


def _mrr(v: str) -> D:
    return (D(v) * 13 / 12).quantize(D("0.01"))


# ── Composição do MRR ────────────────────────────────────────────────────────────────────────────────

def test_a_lista_da_carteira_soma_o_mrr_do_card(cliente):
    card = cliente.get("/api/mrr", params={"hoje": HOJE}, headers=ADMIN).json()
    carteira = cliente.get("/api/mrr/carteira", params={"hoje": HOJE}, headers=ADMIN).json()
    assert D(carteira["mrr"]) == D(card["atual"]["valor"]) == _mrr("12000") + _mrr("2400") + _mrr("3600")
    assert sum(D(g["mrr"]) for g in carteira["itens"]) == D(carteira["mrr"])
    assert carteira["contratos"] == 3 and carteira["grupos"] == 2
    assert [g["grupo"] for g in carteira["itens"]][:2] == ["Grupo Alfa", "Beta"]
    alfa = carteira["itens"][0]
    assert alfa["contratos"] == 2 and [i["escopo"] for i in alfa["itens"]] == ["Contábil", "BPO Financeiro"]


def test_esperado_e_a_parcela_sem_o_13_avos_e_conta_quem_saiu_no_mes(cliente):
    carteira = cliente.get("/api/mrr/carteira", params={"hoje": HOJE, "competencia": "2026-10"}, headers=ADMIN).json()
    por_nome = {g["grupo"]: g for g in carteira["itens"]}
    assert D(por_nome["Grupo Alfa"]["esperado"]) == D("14400")
    assert D(por_nome["Beta"]["esperado"]) == D("3600")  # o suspenso não fatura
    # Gama saiu em 15/10 (10/10/2026: o mês da saída tem parcela e a 13ª proporcional, 10 avos: jan a out).
    assert D(por_nome["Gama"]["esperado"]) == D("6000") + D("5000") and D(por_nome["Gama"]["mrr"]) == 0
    setembro = cliente.get("/api/mrr/carteira", params={"hoje": HOJE, "competencia": "2026-09"}, headers=ADMIN).json()
    por_nome = {g["grupo"]: g for g in setembro["itens"]}
    assert D(por_nome["Gama"]["esperado"]) == D("6000") and D(por_nome["Gama"]["mrr"]) == 0
    assert D(por_nome["Beta"]["esperado"]) == D("3000")  # antes do reajuste de 10/10
    bpo = [i for i in por_nome["Grupo Alfa"]["itens"] if i["escopo"] == "BPO Financeiro"][0]
    assert D(bpo["esperado"]) == 0 and bpo["parte"] == "somado"  # está no MRR de hoje; assinado em 05/10


def test_dezembro_espera_a_13a_parcela():
    assert regra.esperado_no_mes(D("1000"), date(2026, 12, 1)) == D("2000")
    assert regra.esperado_no_mes(D("1000"), date(2026, 11, 1)) == D("1000")


def test_13a_em_avos_como_no_13o_salario():
    # Entrou em 05/03/2026: março conta (27 dias), então 10 avos em dezembro.
    assert regra.avos(2026, date(2026, 3, 5), None) == 10
    assert regra.esperado_no_mes(D("1200"), date(2026, 12, 1), inicio=date(2026, 3, 5)) == D("2200")
    # Entrou em 20/03: março tem 12 dias, não conta.
    assert regra.avos(2026, date(2026, 3, 20), None) == 9


def test_na_saida_paga_a_13a_proporcional_junto_com_a_ultima_parcela():
    # Exemplos aprovados por Eduardo (10/10/2026): parcela de R$ 3.600.
    assert regra.esperado_no_mes(D("3600"), date(2026, 12, 1), saida=date(2026, 12, 9)) == D("3600") + D("3300")  # 11 avos
    assert regra.esperado_no_mes(D("3600"), date(2026, 10, 1), saida=date(2026, 10, 31)) == D("3600") + D("3000")  # 10 avos
    # Quem saiu antes de dezembro não tem 13ª em dezembro.
    assert regra.decima_terceira(D("3600"), date(2026, 12, 1), None, date(2026, 10, 31)) == 0


def test_sem_planilha_o_mes_fica_sem_registro(cliente):
    carteira = cliente.get("/api/mrr/carteira", params={"hoje": HOJE}, headers=ADMIN).json()
    assert carteira["mes_importado"] is False and D(carteira["recebido"]) == 0
    assert {g["situacao"] for g in carteira["itens"]} == {"Sem registro"}


def test_a_lista_do_movimento_soma_cada_linha(cliente):
    mrr = cliente.get("/api/mrr", params={"hoje": HOJE, "de": "2026-10-01"}, headers=ADMIN).json()["movimento"]
    itens = cliente.get("/api/mrr/movimento", params={"hoje": HOJE, "de": "2026-10-01"}, headers=ADMIN).json()["itens"]
    for categoria in ("novo", "expansao", "reajuste", "contracao", "churn_cliente", "churn_criterio"):
        assert sum((D(i["valor"]) for i in itens if i["categoria"] == categoria), D("0")) == D(mrr[categoria]), categoria
    churn = [i for i in itens if i["categoria"] == "churn_cliente"]
    assert churn[0]["grupo"] == "Gama" and churn[0]["data"] == "2026-10-15"


# ── Importação ───────────────────────────────────────────────────────────────────────────────────────

PLANILHA = [
    ["CNPJ", "Grupo", "Competência", "Valor recebido", "Data do recebimento"],
    ["11.222.333/0001-81", "", "10/2026", 10000, "10/10/2026"],
    ["", "beta servicos", "out/2026", "R$ 3.600,00", "12/10/2026"],
    ["", "Desconhecida", "10/2026", 500, ""],
    ["99999999000199", "", "10/2026", 100, ""],
    ["", "Grupo Alfa", "13/2026", 100, ""],
]


def test_previa_nao_grava_e_mostra_o_que_ficou_de_fora(cliente, engine):
    r = cliente.post("/api/recebimentos/previa", json={"arquivo": "out.xlsx", "conteudo_base64": _xlsx(PLANILHA)}, headers=ADMIN)
    assert r.status_code == 200, r.text
    p = r.json()
    assert p["reconhecidas"] == 2 and D(p["total"]) == D("13600")
    assert [(m["competencia"], m["substitui"]) for m in p["meses"]] == [("2026-10", None)]
    assert [(x["linha"], x["motivo"]) for x in p["problemas"]] == [
        (4, "nome não encontrado entre os grupos e empresas do CRM"),
        (5, "CNPJ não cadastrado em nenhuma empresa do CRM"),
        (6, "competência ilegível (use mês/ano, como 10/2026)"),
    ]
    with Session(engine) as s:
        assert s.query(Recebimento).count() == 0


def test_importar_mostra_recebido_e_situacao_e_reimportar_substitui(cliente, engine):
    r = cliente.post("/api/recebimentos/importar", json={"arquivo": "out.xlsx", "conteudo_base64": _xlsx(PLANILHA)}, headers=ADMIN)
    assert r.status_code == 200, r.text
    assert r.json()["linhas"] == 2 and r.json()["fora"] == 3 and r.json()["importado_por"]
    carteira = cliente.get("/api/mrr/carteira", params={"hoje": HOJE}, headers=ADMIN).json()
    por_nome = {g["grupo"]: g for g in carteira["itens"]}
    assert carteira["mes_importado"] is True and D(carteira["recebido"]) == D("13600")
    assert por_nome["Grupo Alfa"]["situacao"] == "Parcial" and por_nome["Beta"]["situacao"] == "Em dia"
    assert carteira["grupos_registrados"] == 2

    nova = _xlsx([["Grupo", "Competência", "Valor"], ["Beta", "2026-10", 1000]])
    previa = cliente.post("/api/recebimentos/previa", json={"arquivo": "out2.csv.xlsx", "conteudo_base64": nova}, headers=ADMIN).json()
    assert D(previa["meses"][0]["substitui"]) == D("13600")
    cliente.post("/api/recebimentos/importar", json={"arquivo": "out2.xlsx", "conteudo_base64": nova}, headers=ADMIN)
    carteira = cliente.get("/api/mrr/carteira", params={"hoje": HOJE}, headers=ADMIN).json()
    por_nome = {g["grupo"]: g for g in carteira["itens"]}
    assert por_nome["Grupo Alfa"]["situacao"] == "Em aberto" and por_nome["Beta"]["situacao"] == "Parcial"
    with Session(engine) as s:
        assert s.query(Recebimento).count() == 3  # as antigas ficam, desligadas
    historico = cliente.get("/api/recebimentos/importacoes", headers=ADMIN).json()
    assert [h["arquivo"] for h in historico] == ["out2.xlsx", "out.xlsx"]


def test_recebido_de_quem_nao_tem_parcela_fica_fora_da_conta(cliente):
    gama = _xlsx([["Grupo", "Competência", "Valor"], ["Gama", "11/2026", 6000]])  # Gama saiu em 15/10
    cliente.post("/api/recebimentos/importar", json={"arquivo": "g.xlsx", "conteudo_base64": gama}, headers=ADMIN)
    carteira = cliente.get("/api/mrr/carteira", params={"hoje": HOJE, "competencia": "2026-11"}, headers=ADMIN).json()
    assert D(carteira["recebido"]) == 0 and D(carteira["recebido_fora"]) == D("6000")
    assert {g["grupo"]: g["situacao"] for g in carteira["itens"]}["Gama"] is None


def test_csv_com_ponto_e_virgula():
    conteudo = "CNPJ;Competência;Valor recebido;Data\n11222333000181;10/2026;1.234,56;05/10/2026\n".encode("latin-1")
    leitura = regra.ler_planilha(conteudo, "rec.csv")
    assert leitura.linhas[0].valor == D("1234.56") and leitura.linhas[0].data == date(2026, 10, 5)


def test_planilha_sem_cabecalho_diz_o_que_falta(cliente):
    r = cliente.post("/api/recebimentos/previa", json={"arquivo": "x.xlsx", "conteudo_base64": _xlsx([["a", "b"], [1, 2]])}, headers=ADMIN)
    assert r.status_code == 422 and "Competência" in r.json()["detail"]


def test_so_o_administrador_importa(cliente):
    corpo = {"arquivo": "out.xlsx", "conteudo_base64": _xlsx(PLANILHA)}
    assert cliente.post("/api/recebimentos/previa", json=corpo, headers=KARINE).status_code == 403
    assert cliente.post("/api/recebimentos/importar", json=corpo, headers=KARINE).status_code == 403
    assert cliente.get("/api/mrr/carteira", params={"hoje": HOJE}, headers=KARINE).status_code == 200


def test_fusao_leva_os_recebimentos(cliente, engine):
    corpo = _xlsx([["Grupo", "Competência", "Valor"], ["Beta", "10/2026", 3600]])
    cliente.post("/api/recebimentos/importar", json={"arquivo": "b.xlsx", "conteudo_base64": corpo}, headers=ADMIN)
    from crm.db.grupos import fundir_grupos

    with Session(engine) as s:
        alfa = s.query(GrupoEconomico).filter_by(nome="Grupo Alfa").one()
        beta = s.query(GrupoEconomico).filter_by(nome="Beta").one()
        fundir_grupos(s, alfa, beta)
        s.commit()
        assert {r.grupo_id for r in s.query(Recebimento)} == {alfa.id}


# ── Meta em 13 parcelas ──────────────────────────────────────────────────────────────────────────────

def test_meta_de_origem_vira_13_parcelas(engine):
    with Session(engine) as s:
        s.add(MetaDeIndicador(chave="mrr", meta=D("400000"), alerta=D("200000")))
        s.commit()
        assert "convertida" in meta_de_mrr_em_13_parcelas.executar(s)
        linha = s.get(MetaDeIndicador, "mrr")
        assert (linha.meta, linha.alerta) == (D("433333.33"), D("216666.67"))
        assert "nada a converter" in meta_de_mrr_em_13_parcelas.executar(s)


def test_meta_mudada_na_tela_fica_como_esta(engine):
    with Session(engine) as s:
        s.add(MetaDeIndicador(chave="mrr", meta=D("450000"), alerta=D("250000")))
        s.commit()
        assert "fica como está" in meta_de_mrr_em_13_parcelas.executar(s)
        assert s.get(MetaDeIndicador, "mrr").meta == D("450000")


def test_sem_linha_vale_o_padrao_ja_em_13_parcelas(engine):
    with Session(engine) as s:
        assert "padrão do código" in meta_de_mrr_em_13_parcelas.executar(s)


# ── Parcelas no ano por contrato (11/10/2026) ───────────────────────────────────────────────────────

def test_contrato_de_12_parcelas_tem_mrr_igual_a_parcela_e_sem_13a(cliente, engine):
    """A Menino da Ilha está no formato antigo, sem 13ª: MRR = parcela; em dezembro, o esperado não soma a 13ª."""
    with Session(engine) as s:
        b1 = s.scalars(sa.select(Contrato).where(Contrato.escopo == "Contábil", Contrato.preco_mensal == D("3600"))).one()
        b1.parcelas_no_ano = 12
        s.commit()
    carteira = cliente.get("/api/mrr/carteira", params={"hoje": HOJE}, headers=ADMIN).json()
    beta = next(g for g in carteira["itens"] if g["grupo"] == "Beta")
    assert D(beta["mrr"]) == D("3600.00")  # e não 3.900 (× 13 ÷ 12)
    assert D(carteira["mrr"]) == _mrr("12000") + _mrr("2400") + D("3600.00")
    dez = cliente.get("/api/mrr/carteira", params={"hoje": "2026-12-20", "competencia": "2026-12"}, headers=ADMIN).json()
    beta = next(g for g in dez["itens"] if g["grupo"] == "Beta")
    assert D(beta["esperado"]) == D("3600")  # sem a 13ª


def test_tarefa_poe_12_parcelas_so_na_menino_da_ilha(engine):
    from crm.manutencao.tarefas import menino_da_ilha_em_12_parcelas

    with Session(engine) as s:
        g = GrupoEconomico(nome="Menino da Ilha")
        s.add(g)
        s.flush()
        bpo = Contrato(grupo_id=g.id, situacao=SituacaoContrato.ATIVO, preco_mensal=D("2000"), escopo="BPO Financeiro")
        s.add(bpo)
        s.commit()
        assert "1 contrato" in menino_da_ilha_em_12_parcelas.executar(s)
        s.commit()
        s.refresh(bpo)
        assert bpo.parcelas_no_ano == 12
        outros = s.scalars(sa.select(Contrato).where(Contrato.id != bpo.id)).all()
        assert all(c.parcelas_no_ano == 13 for c in outros)
        assert "nada a mudar" in menino_da_ilha_em_12_parcelas.executar(s)


# ── Contas a receber do Omie (19/10/2026) ────────────────────────────────────────────────────────────

def _omie() -> str:
    """O formato do contas a receber do Omie: resumo, por mês (sem CNPJ) e a aba Títulos, que é a que vale."""
    from datetime import datetime
    import io

    livro = Workbook()
    resumo = livro.active
    resumo.title = "Resumo"
    resumo.append([None, "Indicador", "Valor"])
    por_mes = livro.create_sheet("Por Mês")
    por_mes.append([None, "Cliente", "Mês", "Faturado", "Recebido", "Vencido"])
    por_mes.append([None, "Beta Serviços", "Out/26", 3600, 3600, 0])
    titulos = livro.create_sheet("Títulos")
    titulos.append(["Contas a Receber 2026 - títulos"])
    titulos.append([])
    titulos.append(["Empresa", "Cliente", "CNPJ/CPF", "Parcela", "Emissão", "Vencimento", "Mês", "Valor do Título",
                    "Recebido", "Em Aberto", "Em Atraso", "Situação", "Dias em Atraso", "Data do Recebimento"])
    def titulo(cliente, doc, emissao, venc, valor, recebido, pago=None):
        titulos.append(["CCE", cliente, doc, "001/013", emissao, venc, venc.month, valor, recebido,
                        valor - recebido, 0, "", None, pago])
    # a competência é o mês da emissão: emitido em 25/10 e vencido em 03/11 é outubro
    titulo("ALFA COMERCIO LTDA", "11.222.333/0001-81", datetime(2026, 10, 25), datetime(2026, 11, 3), 12000, 12000,
           datetime(2026, 11, 3))
    titulo("Beta Serviços", "98.765.432/0001-10", datetime(2026, 10, 1), datetime(2026, 10, 25), 3600, 3600,
           datetime(2026, 10, 25))  # CNPJ fora do cadastro
    titulo("Beta Serviços", "123.456.789-09", datetime(2026, 9, 1), datetime(2026, 9, 25), 3600, 1000,
           datetime(2026, 9, 25))  # CPF: vale o nome
    titulo("Grupo Alfa", "11.222.333/0001-81", datetime(2026, 11, 25), datetime(2026, 12, 3), 12000, 0)  # a vencer
    titulo("Desconhecida SA", None, datetime(2026, 10, 1), datetime(2026, 10, 25), 500, 500, datetime(2026, 10, 25))
    saida = io.BytesIO()
    livro.save(saida)
    return base64.b64encode(saida.getvalue()).decode()


def test_contas_a_receber_do_omie_entra_como_sai(cliente):
    p = cliente.post("/api/recebimentos/previa", json={"arquivo": "omie.xlsx", "conteudo_base64": _omie()}, headers=ADMIN).json()
    assert p["reconhecidas"] == 3 and D(p["total"]) == D("16600")
    assert [(m["competencia"], D(m["total"])) for m in p["meses"]] == [("2026-09", D("1000")), ("2026-10", D("15600"))]
    assert p["sem_recebimento"] == 1
    assert [(x["linha"], x["motivo"]) for x in p["problemas"]] == [(8, "nome não encontrado entre os grupos e empresas do CRM")]

    cliente.post("/api/recebimentos/importar", json={"arquivo": "omie.xlsx", "conteudo_base64": _omie()}, headers=ADMIN)
    carteira = cliente.get("/api/mrr/carteira", params={"hoje": HOJE}, headers=ADMIN).json()
    por_nome = {g["grupo"]: g for g in carteira["itens"]}
    assert D(por_nome["Grupo Alfa"]["recebido"]) == D("12000") and D(por_nome["Beta"]["recebido"]) == D("3600")
    assert por_nome["Beta"]["situacao"] == "Em dia"


def test_cnpj_fora_do_cadastro_sem_nome_continua_de_fora():
    leitura = regra.ler_planilha("CNPJ;Competência;Valor\n123;10/2026;10\n".encode(), "r.csv")
    assert [p.motivo for p in leitura.problemas] == ["CNPJ sem 14 dígitos"]
