"""Integralidade do cadastro (amostra e lista de campos aprovadas por Eduardo em 13/10/2026): pontos dos campos
devidos e vazios ÷ pontos de todos os devidos, obrigatório 2 e prioritário 1, só o que nasceu desde o corte, e o
legado à parte."""

from __future__ import annotations

from datetime import date
from decimal import Decimal as D

import pytest
from sqlalchemy.orm import Session

from crm.api import integralidade
from crm.api.kpis_lideranca import calcular as calcular_kpis
from crm.db.modelos import Contrato, Empresa, GrupoEconomico, Oportunidade, PessoaContato
from crm.domain.listas import MotivoRecusa, Origem, Situacao, SituacaoContrato, TipoCanal


@pytest.fixture
def cadastro(engine, monkeypatch):
    monkeypatch.setattr(integralidade, "DATA_DE_CORTE", date(2000, 1, 1))  # tudo o que o teste cria é "novo"
    with Session(engine) as s:
        cliente = GrupoEconomico(nome="Cliente X", responsavel_cs="KN")
        prospect = GrupoEconomico(nome="Prospect")
        s.add_all([cliente, prospect])
        s.flush()
        empresa = Empresa(grupo_id=cliente.id, razao_social="Cliente X Ltda", cnpj="11222333000181")
        s.add(empresa)
        s.flush()
        s.add(Contrato(grupo_id=cliente.id, empresa_id=empresa.id, situacao=SituacaoContrato.ATIVO,
                       preco_mensal=D("3000"), data_inicio=date(2026, 9, 1), escopo="BPO Contábil"))
        s.add(Oportunidade(grupo_id=prospect.id, nome="Prospect BPO", servico="BPO Financeiro", situacao=Situacao.PERDIDA,
                           captador="EL", tipo_canal=TipoCanal.SOCIOS, preco_mensal=D("5000"),
                           data_colocacao=date(2026, 9, 2), data_envio_proposta=date(2026, 9, 5)))
        s.add(PessoaContato(nome="Fulana", email="fulana@exemplo.com"))
        s.commit()
    return engine


def test_conta_os_pontos_dos_campos_devidos_com_peso_e_condicao(cadastro):
    with Session(cadastro) as s:
        r = integralidade.calcular(s)
    faltam = {(c.cadastro, c.campo) for c in r.novos if not c.preenchido}
    assert faltam == {
        ("Contrato", "Bruto ou líquido"), ("Contrato", "Signatário"), ("Contrato", "Fim da vigência"),
        ("Contrato", "Data de competência"), ("Contrato", "Início do pagamento"),  # prioritários (28/10/2026)
        ("Cliente", "Classificação (classe e semáforo)"), ("Cliente", "Porte"),
        ("Empresa", "Regime tributário"), ("Empresa", "Endereço (logradouro, município e CEP)"),
        ("Oportunidade", "Motivo de perda"), ("Oportunidade", "Porte"),
        ("Contato", "Vínculo com a empresa"), ("Contato", "Cargo"), ("Contato", "Papel (decisor, influenciador…)"),
    }
    # devidos: contrato 12 + cliente 5 + empresa 4 + oportunidade 15 + contato 6 = 42; vazios: 6 + 3 + 2 + 3 + 4 = 18
    assert sum(c.pontos for c in r.novos) == 42
    assert integralidade.Resultado.pendente_pct(r.novos) == D("42.9")
    # perdida não cobra temperatura nem próxima ação; data de envio já é devida fora de "Enviar proposta"
    campos_da_oportunidade = {c.campo for c in r.novos if c.cadastro == "Oportunidade"}
    assert "Temperatura" not in campos_da_oportunidade and "Data de envio da proposta" in campos_da_oportunidade


def test_quem_responde_e_o_captador_e_o_cs_do_grupo(cadastro):
    with Session(cadastro) as s:
        r = integralidade.calcular(s)
    dono = {c.cadastro: c.responsavel for c in r.novos}
    assert dono == {"Oportunidade": "EL", "Contrato": "KN", "Cliente": "KN", "Empresa": "KN", "Contato": None}


def test_o_legado_fica_a_parte(cadastro):
    with Session(cadastro) as s:
        prospect = s.query(GrupoEconomico).filter_by(nome="Prospect").one()
        s.add(Oportunidade(grupo_id=prospect.id, nome="Antiga", servico="Consultoria", situacao=Situacao.PERDIDA,
                           motivo_recusa=list(MotivoRecusa)[0], origem=Origem.CARGA_2026, chave_origem="antiga|2026"))
        s.commit()
        r = integralidade.calcular(s)
    assert {c.registro for c in r.legado} == {"Prospect · Consultoria"}
    assert integralidade.Resultado.pendente_pct(r.novos) == D("42.9")  # o legado não mexe no número principal


def test_o_card_mostra_novos_em_pct_e_a_lista_por_registro(cadastro):
    with Session(cadastro) as s:
        k = {k.chave: k for k in calcular_kpis(s, "2026-10", date(2026, 10, 20)).kpis}["integralidade"]
    assert k.titulo == "Integralidade do cadastro" and k.valor == D("42.9")
    x = k.extras
    assert (x["novos_pendencias"], x["novos_obrigatorios"], x["novos_pontos_devidos"]) == (14, 4, 42)
    assert x["passado_pendencias"] == 0 and x["passado_retrato"] is None and x["passado_consertadas"] is None
    assert "retrato do Passado ainda não foi guardado" in k.falta[0]
    # Um item por registro, com tudo o que falta nele: o obrigatório primeiro e marcado com "!".
    contrato = next(i for i in k.itens if i.grupo == "Contrato · Cliente X · BPO Contábil")
    assert contrato.detalhe == ("! Bruto ou líquido, Signatário, Fim da vigência, Data de competência, "
                                "Início do pagamento")
    assert (contrato.corte, contrato.cadastro, contrato.responsavel) == ("Novos", "Contrato", "KN")
    assert [(f.campo, f.obrigatorio) for f in contrato.faltam][0] == ("Bruto ou líquido", True)
    assert len(k.itens) == 5  # contrato, cliente, empresa, oportunidade e contato
    assert [i.responsavel for i in k.itens][-1] == "Sem responsável"
    assert all(i.valor == 0 for i in k.itens)


def test_passado_com_retrato_conta_as_consertadas(engine, monkeypatch):
    monkeypatch.setattr(integralidade, "DATA_DE_CORTE", date(2999, 1, 1))  # tudo é Passado
    with Session(engine) as s:
        g = GrupoEconomico(nome="Antigo")
        s.add(g)
        s.flush()
        e1, e2 = Empresa(grupo_id=g.id, razao_social="Sem CNPJ 1"), Empresa(grupo_id=g.id, razao_social="Sem CNPJ 2")
        s.add_all([e1, e2])
        s.commit()
        retrato = integralidade.guardar_retrato(s, date(2026, 10, 15))
        assert (retrato.pendencias, retrato.obrigatorios) == (6, 2)  # 2 empresas × (CNPJ, regime, endereço)
        assert integralidade.guardar_retrato(s, date(2026, 10, 16)) is None  # uma vez só
        e1.cnpj, e1.regime_tributario = "11222333000181", "Simples Nacional"
        s.commit()
        k = {k.chave: k for k in calcular_kpis(s, "2026-10", date(2026, 10, 20)).kpis}["integralidade"]
    x = k.extras
    assert (x["passado_pendencias"], x["passado_obrigatorios"], x["passado_retrato"], x["passado_consertadas"]) == (4, 1, 6, 2)
    assert x["passado_empresa"] == 4 and x["passado_retrato_em"] == "2026-10-15" and k.falta == []
    assert k.valor is None  # nenhum registro Novo: sem dado
    assert {i.corte for i in k.itens} == {"Passado"}


def test_tarefa_do_redeploy_guarda_o_retrato_uma_vez(engine):
    from crm.manutencao.tarefas import retrato_da_integralidade
    with Session(engine) as s:
        primeira = retrato_da_integralidade.executar(s)
        segunda = retrato_da_integralidade.executar(s)
    assert primeira.startswith("retrato do Passado em") and segunda.startswith("nada a fazer")


def test_planilha_das_pendencias(cadastro):
    from fastapi.testclient import TestClient
    from io import BytesIO
    from openpyxl import load_workbook
    from sqlalchemy.orm import sessionmaker
    from crm.api.app import criar_app

    fabrica = sessionmaker(bind=cadastro, expire_on_commit=False, future=True)
    with TestClient(criar_app(fabrica)) as c:
        tudo = c.get("/api/inteligencia/integralidade/planilha")
        passado = c.get("/api/inteligencia/integralidade/planilha", params={"corte": "Passado"})
    assert tudo.status_code == 200 and "pendencias-de-cadastro-" in tudo.headers["content-disposition"]
    ws = load_workbook(BytesIO(tudo.content))["Pendências"]
    assert [c.value for c in ws[1]] == ["Corte", "Responsável", "Cadastro", "Registro", "Campo em falta", "Nível"]
    assert ws.max_row == 15  # cabeçalho + 14 campos em falta
    assert load_workbook(BytesIO(passado.content))["Pendências"].max_row == 1


def test_todo_campo_tem_chave_e_o_registro_leva_o_grupo(cadastro):
    """A tela acha o campo pela chave e abre o card pelo grupo (18/10/2026)."""
    with Session(cadastro) as s:
        r = integralidade.calcular(s)
        cliente = s.query(GrupoEconomico).filter_by(nome="Cliente X").one()
    assert {c.chave for c in r.campos} <= set(integralidade.CHAVES.values())  # nenhum campo sem chave (KeyError)
    assert len(set(integralidade.CHAVES)) == len(integralidade.CHAVES)
    por = {(c.cadastro, c.campo): c for c in r.campos}
    contrato = por[("Contrato", "Bruto ou líquido")]
    assert (contrato.chave, contrato.grupo_id, contrato.contrato_id) == ("base_do_valor", cliente.id, contrato.registro_id)
    empresa = por[("Empresa", "Regime tributário")]
    assert (empresa.chave, empresa.grupo_id, empresa.empresa_id) == ("regime_tributario", cliente.id, empresa.registro_id)
    assert por[("Contato", "Cargo")].grupo_id is None  # sem vínculo: sem grupo


def test_a_composicao_leva_a_chave_e_o_registro(cadastro):
    with Session(cadastro) as s:
        k = {k.chave: k for k in calcular_kpis(s, "2026-10", date(2026, 10, 20)).kpis}["integralidade"]
    contrato = next(i for i in k.itens if i.cadastro == "Contrato")
    assert [f.chave for f in contrato.faltam] == ["base_do_valor", "signatario", "data_fim", "data_competencia",
                                                    "data_inicio_pagamento"]
    assert contrato.grupo_id is not None and contrato.registro_id == contrato.contrato_id


def test_lacunas_do_grupo_e_do_registro_pela_api(cadastro):
    from fastapi.testclient import TestClient
    from sqlalchemy.orm import sessionmaker
    from crm.api.app import criar_app

    with Session(cadastro) as s:
        cliente = s.query(GrupoEconomico).filter_by(nome="Cliente X").one()
        oportunidade = s.query(Oportunidade).one()
    fabrica = sessionmaker(bind=cadastro, expire_on_commit=False, future=True)
    with TestClient(criar_app(fabrica)) as c:
        do_grupo = c.get("/api/lacunas", params={"grupo_id": cliente.id}).json()
        da_oportunidade = c.get("/api/lacunas", params={"cadastro": "Oportunidade", "registro_id": oportunidade.id}).json()
        sem_filtro = c.get("/api/lacunas")
        grupo = c.get(f"/api/grupos/{cliente.id}")
        nenhum = c.get("/api/grupos/999999")
        sugestoes = c.get("/api/grupos/sugestoes-de-fusao")
    assert [(l["cadastro"], l["chave"]) for l in do_grupo] == [
        ("Contrato", "base_do_valor"), ("Contrato", "signatario"), ("Contrato", "data_fim"),
        ("Contrato", "data_competencia"), ("Contrato", "data_inicio_pagamento"),
        ("Cliente", "classificacao"), ("Cliente", "porte"),
        ("Empresa", "regime_tributario"), ("Empresa", "endereco"),
    ]
    assert [(l["chave"], l["obrigatorio"]) for l in da_oportunidade] == [("motivo_recusa", True), ("porte", False)]
    assert sem_filtro.status_code == 422
    assert grupo.status_code == 200 and grupo.json()["nome"] == "Cliente X"
    assert nenhum.status_code == 404 and sugestoes.status_code == 200  # a rota do id não engole as outras


def test_contrato_recorrente_nao_cobra_o_fim_da_vigencia(cadastro):
    """Recorrente vale por prazo indeterminado (Eduardo, 23/10/2026); o de projeto continua cobrando o fim."""
    with Session(cadastro) as s:
        grupo = s.query(GrupoEconomico).filter_by(nome="Cliente X").one()
        empresa = s.query(Empresa).filter_by(grupo_id=grupo.id).one()
        s.query(Contrato).update({"recorrente": True})
        s.add(Contrato(grupo_id=grupo.id, empresa_id=empresa.id, situacao=SituacaoContrato.ATIVO, recorrente=False,
                       preco_anual=D("24000"), data_inicio=date(2026, 9, 1), escopo="Projeto de valuation"))
        s.commit()
        r = integralidade.calcular(s)
    fim = [c.registro for c in r.novos if c.cadastro == "Contrato" and c.campo == "Fim da vigência"]
    assert fim == ["Cliente X · Projeto de valuation"]
